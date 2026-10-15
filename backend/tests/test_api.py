import io
import zipfile

import docx
import pytest
from openpyxl import Workbook

from .conftest import make_pdf


def test_health(client):
    assert client.get("/api/health").json()["status"] == "ok"


def test_ai_status_without_provider_is_not_configured(client):
    r = client.get("/api/ai/status").json()
    assert r["configured"] is False


def test_appraisal_endpoint(client):
    r = client.post("/api/finance/appraisal", json={"cash_flows": [-1000, 500, 500, 500], "discount_rate": 0.1}).json()
    assert r["npv"] == pytest.approx(243.426, abs=1e-3)
    assert r["irr"]["status"] == "unique"
    assert r["timeline"][0]["present_value"] == -1000
    assert r["timeline"][-1]["cumulative_discounted"] == pytest.approx(r["npv"])


def test_appraisal_validation_error_is_readable(client):
    r = client.post("/api/finance/appraisal", json={"cash_flows": [-1000], "discount_rate": 0.1})
    assert r.status_code == 422 and "cash_flows" in r.json()["detail"]


def test_dcf_invalid_growth_returns_422(client):
    body = {"base_revenue": 100, "revenue_growth": [0.1], "ebit_margin": [0.2], "tax_rate": 0.25, "da_pct_revenue": 0,
            "capex_pct_revenue": 0, "nwc_pct_revenue": 0, "wacc": 0.05, "terminal_growth": 0.06}
    r = client.post("/api/finance/dcf", json=body)
    assert r.status_code == 422 and "terminal growth" in r.json()["detail"]


def test_dcf_sensitivity_endpoint(client):
    model = {"base_revenue": 100, "revenue_growth": [0.1] * 3, "ebit_margin": [0.2], "tax_rate": 0.25, "da_pct_revenue": 0.03,
             "capex_pct_revenue": 0.04, "nwc_pct_revenue": 0.1, "wacc": 0.1, "terminal_growth": 0.02}
    r = client.post("/api/finance/dcf/sensitivity", json={"model": model, "x_variable": "wacc", "x_values": [0.02, 0.1],
                                                           "y_variable": "terminal_growth", "y_values": [0.02]}).json()
    assert r["grid"][0][0]["value"] is None and r["grid"][0][0]["error"]
    assert r["grid"][0][1]["value"] == pytest.approx(r["base_output"])


def test_variance_endpoint_zero_budget(client):
    r = client.post("/api/finance/variance", json={"lines": [
        {"name": "Sales", "kind": "revenue", "budget": 100, "actual": 90},
        {"name": "New cost", "kind": "expense", "budget": 0, "actual": 5}]}).json()
    assert r["lines"][0]["status"] == "adverse"
    assert r["lines"][1]["variance_pct"] is None
    assert r["totals"]["profit"]["variance"] == pytest.approx(-15)


def test_ratios_endpoint_rejects_unknown_items(client):
    r = client.post("/api/finance/ratios", json={"data": {"made_up": 1}})
    assert r.status_code == 422


def test_dashboard_sample_is_flagged(client):
    r = client.get("/api/dashboard/sample").json()
    assert r["is_sample"] is True and "sample" in r["company"].lower()


# ---------------------------------------------------------------- uploads

def xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    ws.append(["Line item", "FY2024", "FY2025"])
    ws.append(["Revenue", 1000, 1200])
    ws.append(["Cost of goods sold", 600, 700])
    ws.append(["Net income", 120, 150])
    ws["D1"] = "Growth"
    ws["D2"] = "=C2/B2-1"  # formula with no cached value (openpyxl does not calculate)
    ws2 = wb.create_sheet("Budget")
    ws2.append(["Account", "Budget", "Actual"])
    ws2.append(["Sales revenue", 500, 540])
    ws2.append(["Salaries", 200, 230])
    ws2.append(["Rent", 50, 50])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, name, data):
    r = client.post("/api/uploads", files=[("files", (name, data))])
    assert r.status_code == 200
    return r.json()["results"][0]


def test_upload_xlsx_detects_formulas_mapping_and_variance(client):
    res = upload(client, "model.xlsx", xlsx_bytes())
    assert res["ok"], res
    up = res["upload"]
    modules = {r["module"] for r in up["recommendations"]}
    assert "Variance Analysis" in modules and "Ratio Analysis" in modules
    mapping = {m["item"]: m for m in up["statement_mapping"]}
    assert mapping["revenue"]["value"] == 1200 and mapping["revenue"]["prior_value"] == 1000
    assert mapping["revenue"]["source"].endswith("C2")
    lines = up["variance_lines"]["lines"]
    assert {l["name"] for l in lines} == {"Sales revenue", "Salaries", "Rent"}
    assert next(l for l in lines if l["name"] == "Sales revenue")["kind"] == "revenue"

    full = client.get(f"/api/uploads/{up['id']}").json()
    pl = full["analysis"]["sheets"][0]
    assert pl["formula_count"] == 1 and pl["formulas_without_cached_value"] == 1
    assert any("recalculated" in w for w in pl["warnings"])

    series = client.get("/api/uploads/series").json()["series"]
    assert any(s["label"] == "Revenue" and s["refs"] == ["B2", "C2"] for s in series) is False  # rows need 3+ values
    assert client.delete(f"/api/uploads/{up['id']}").status_code == 200
    assert client.get(f"/api/uploads/{up['id']}").status_code == 404


def test_upload_csv_with_text_numbers(client):
    data = b"Month,Revenue\nJan,\"1,000\"\nFeb,1100\nMar,(50)\nApr,1300\n"
    up = upload(client, "rev.csv", data)["upload"]
    full = client.get(f"/api/uploads/{up['id']}").json()
    s = full["analysis"]["sheets"][0]["series"][0]
    assert s["values"] == [1000, 1100, -50, 1300]
    assert s["refs"][0] == "row 2, col 2"
    assert any("stored as text" in w for w in full["analysis"]["sheets"][0]["warnings"])


def test_upload_docx_and_pdf(client):
    d = docx.Document()
    d.add_paragraph("Annual report FY2025. All figures in ₹ lakhs.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Value"
    t.cell(1, 0).text, t.cell(1, 1).text = "Revenue", "1,000"
    buf = io.BytesIO()
    d.save(buf)
    up = upload(client, "report.docx", buf.getvalue())["upload"]
    full = client.get(f"/api/uploads/{up['id']}").json()["analysis"]
    assert "INR" in full["context"]["currencies"] and "lakhs" in full["context"]["units"]
    assert "FY2025" in full["context"]["periods"]

    pdf = upload(client, "statement.pdf", make_pdf(["Balance Sheet FY2025", "Total assets 5000"]))
    assert pdf["ok"], pdf
    full = client.get(f"/api/uploads/{pdf['upload']['id']}").json()["analysis"]
    assert full["page_count"] == 1 and "Total assets" in full["text"]


def test_upload_python_is_not_executed_and_syntax_checked(client, tmp_path):
    marker = tmp_path / "pwned"
    code = f"import os\nopen({str(marker)!r}, 'w').write('x')\ndef f(:\n".encode()
    up = upload(client, "evil.py", code)["upload"]
    full = client.get(f"/api/uploads/{up['id']}").json()["analysis"]
    assert not marker.exists()
    assert full["outline"]["syntax_ok"] is False
    assert up["recommendations"][0]["module"] == "Data & Code Lab"


def test_upload_rejects_mismatched_content_and_unknown_type(client):
    assert upload(client, "fake.xlsx", b"not a zip")["ok"] is False
    assert upload(client, "fake.pdf", b"hello")["ok"] is False
    assert upload(client, "virus.exe", b"MZ....")["ok"] is False


def test_upload_size_limit(client, monkeypatch):
    monkeypatch.setenv("FIS_MAX_UPLOAD_MB", "1")
    from app.config import get_settings
    get_settings.cache_clear()
    r = upload(client, "big.csv", b"a,b\n" + b"1,2\n" * 400_000)
    assert r["ok"] is False and "limit" in r["error"]


def test_zip_traversal_blocked_and_members_parsed(client):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("../../etc/passwd", "root:x")
        z.writestr("data/sales.csv", "Month,Sales\nJan,1\nFeb,2\nMar,3\n")
        z.writestr("query.sql", "SELECT * FROM t;")
    up = upload(client, "bundle.zip", buf.getvalue())["upload"]
    full = client.get(f"/api/uploads/{up['id']}").json()["analysis"]
    blocked = [m for m in full["members"] if m["blocked"]]
    assert blocked and blocked[0]["blocked"] == "path traversal"
    assert {c["path"] for c in full["children"]} == {"data/sales.csv", "query.sql"}


def test_zip_bomb_rejected(client):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("big.txt", b"0" * 50_000_000)
    r = upload(client, "bomb.zip", buf.getvalue())
    assert r["ok"] is False and "zip bomb" in r["error"]


def test_invalid_upload_id_cannot_traverse(client):
    assert client.get("/api/uploads/..%2F..%2Fetc%2Fpasswd").status_code in (400, 404)
    assert client.get("/api/uploads/" + "z" * 32).status_code == 400


def test_malformed_xlsx_reports_failure_not_crash(client):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("random.txt", "x")
    r = upload(client, "broken.xlsx", buf.getvalue())
    assert r["ok"] and r["upload"]["status"] == "failed"


def test_break_even_endpoint(client):
    model = {"base_revenue": 100, "revenue_growth": [0.05] * 3, "ebit_margin": [0.2], "tax_rate": 0.25, "da_pct_revenue": 0,
             "capex_pct_revenue": 0, "nwc_pct_revenue": 0, "wacc": 0.1, "terminal_growth": 0.02}
    ev12 = client.post("/api/finance/dcf", json={**model, "wacc": 0.12}).json()["enterprise_value"]
    r = client.post("/api/finance/dcf/break-even", json={"model": model, "variable": "wacc", "target": ev12, "low": 0.05, "high": 0.3}).json()
    assert r["value"] == pytest.approx(0.12, abs=1e-6)
    r2 = client.post("/api/finance/dcf/break-even", json={"model": model, "variable": "wacc", "target": -1, "low": 0.05, "high": 0.3}).json()
    assert r2["value"] is None and r2["message"]


# ---------------------------------------------------------------- Phase 3 endpoints

def test_budget_endpoint(client):
    r = client.post("/api/finance/budget", json={"start_month": 4, "start_year": 2026, "lines": [
        {"name": "Sales", "kind": "revenue", "method": "annual", "annual": 1200},
        {"name": "Rent", "kind": "expense", "method": "incremental", "prior_actual": 100, "change_pct": 20},
    ]}).json()
    assert r["months"][0] == "Apr-26"
    assert r["totals_annual"]["operating_result"] == pytest.approx(1080)
    assert sum(r["lines"][0]["monthly"]) == pytest.approx(1200)


def test_budget_endpoint_missing_input(client):
    r = client.post("/api/finance/budget", json={"lines": [{"name": "Rent", "kind": "expense", "method": "annual"}]})
    assert r.status_code == 422 and "annual amount" in r.json()["detail"]


def test_statements_endpoint(client):
    r = client.post("/api/finance/statements", json={"periods": ["FY24", "FY25"], "lines": [
        {"label": "Revenue", "statement": "income", "values": [100, 120], "key": "revenue"},
        {"label": "Total assets", "statement": "balance", "values": [200, 210], "key": "total_assets"},
        {"label": "Total liabilities", "statement": "balance", "values": [120, 100], "key": "total_liabilities"},
        {"label": "Equity", "statement": "balance", "values": [80, 100], "key": "total_equity"},
    ]}).json()
    assert r["lines"][0]["changes"][1]["pct"] == pytest.approx(20)
    statuses = [c["status"] for c in r["checks"]]
    assert statuses == ["pass", "fail"]


def test_cash_budget_and_wc_endpoints(client):
    r = client.post("/api/finance/cash-budget", json={"periods": ["M1", "M2"], "opening_cash": 100, "minimum_cash": 10,
                                                     "receipts": [{"name": "Sales", "values": [10, 10]}],
                                                     "payments": [{"name": "Costs", "values": [60, 60]}]}).json()
    assert r["periods"][1]["financing"] == pytest.approx(10)
    assert r["runway"]["months"] == pytest.approx(2)
    w = client.post("/api/finance/working-capital", json={"revenue": 3650, "cogs": 1825, "receivables": 500, "inventory": 300,
                                                          "payables": 150, "target_dso": 40, "target_dio": 50, "target_dpo": 40}).json()
    assert w["cash_released"] == pytest.approx(200)


def test_upload_statements_import_reorders_descending_years(client):
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"
    ws.append(["Particulars", "FY2025", "FY2024"])
    ws.append(["Total assets", 1300, 1100])
    ws.append(["Total liabilities", 600, 520])
    ws.append(["Total equity", 700, 580])
    buf = io.BytesIO()
    wb.save(buf)
    up = upload(client, "bs.xlsx", buf.getvalue())["upload"]
    st = up["statements"][0]
    assert st["statement"] == "balance" and st["reordered"] is True
    assert st["periods"] == ["FY2024", "FY2025"]
    ta = next(l for l in st["lines"] if l["key"] == "total_assets")
    assert ta["values"] == [1100, 1300] and ta["source"].endswith("C2–B2")
    assert any(r["module"] == "Financial Statements" for r in up["recommendations"])


def test_period_label_detection():
    from app.files.routing import is_period_label
    for ok in ["FY2025", "FY 24", "2023", "Q1 2024", "Q3FY25", "Mar-25", "March 2025", "H1 2024"]:
        assert is_period_label(ok), ok
    for bad in ["Budget", "Actual", "Revenue", "Notes", "Column C"]:
        assert not is_period_label(bad), bad


def test_budget_sheet_not_offered_as_statement(client):
    up = upload(client, "model.xlsx", xlsx_bytes())["upload"]
    assert [s["sheet"] for s in up["statements"]] == ["P&L"]
