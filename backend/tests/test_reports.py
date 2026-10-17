"""Report exports: the Excel formulas must reproduce the engine's numbers."""

import io

import docx
import pytest
from openpyxl import load_workbook

from app.files.formula_eval import WorkbookEvaluator
from app.files.screener import parse_screener
from app.finance.dcf import DCFInputs, run_dcf
from app.finance.profile import build
from app.reports.builder import Line, ReportInput, build_excel, build_word

from .screener_fixture import screener_bytes

DCF = {"base_revenue": 1000, "revenue_growth": [0.12, 0.1, 0.08, 0.06, 0.05], "ebit_margin": [0.2], "tax_rate": 0.25,
       "da_pct_revenue": 0.04, "capex_pct_revenue": 0.05, "nwc_pct_revenue": 0.1, "wacc": 0.11, "terminal_growth": 0.04,
       "terminal_method": "gordon", "exit_multiple": None, "debt": 300, "cash": 120, "shares_outstanding": 50, "mid_year": False}


def report_input():
    r = parse_screener(screener_bytes())
    lines = [Line(l["label"], l["statement"], l["key"], l["values"], l["source"]) for l in r["lines"]]
    return ReportInput("Tata Consultancy Services Ltd", r["periods"], lines, units="crores", market=r["market"], dcf=DCF,
                       memo="# Summary\n- Strong cash conversion", prepared_by="Analyst")


def evaluator(data: bytes) -> WorkbookEvaluator:
    wb = load_workbook(io.BytesIO(data))
    values, formulas = {}, {}
    for ws in wb.worksheets:
        grid, f = [], {}
        for r, row in enumerate(ws.iter_rows()):
            vals = []
            for c, cell in enumerate(row):
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    f[(r, c)] = v
                    v = None
                vals.append(v)
            grid.append(vals)
        values[ws.title], formulas[ws.title] = grid, f
    return WorkbookEvaluator(values, formulas)


def find(ev, sheet, label, col=1):
    grid = ev.values[sheet]
    r = next(i for i, row in enumerate(grid) if row and row[0] == label)
    return ev.cell(sheet, r, col)


def test_excel_dcf_formulas_match_engine():
    data = build_excel(report_input())
    ev = evaluator(data)
    engine = run_dcf(DCFInputs(**DCF))
    assert find(ev, "DCF", "Enterprise value") == pytest.approx(engine.enterprise_value)
    assert find(ev, "DCF", "Equity value") == pytest.approx(engine.equity_value)
    assert find(ev, "DCF", "Value per share") == pytest.approx(engine.value_per_share)
    assert not ev.unsupported


def test_excel_ratio_formulas_match_profile():
    inp = report_input()
    ev = evaluator(build_excel(inp))
    items = {l.key: l.values for l in inp.lines if l.key}
    items.update(inp.market)
    prof = build(inp.periods, items)
    last = len(inp.periods)  # column index of the latest period on the Ratios sheet
    assert find(ev, "Ratios", "Net margin", last) == pytest.approx(prof.series["net_margin"][-1])
    assert find(ev, "Ratios", "Return on equity (avg)", last) == pytest.approx(prof.series["roe"][-1])
    assert find(ev, "Ratios", "P/E", last) == pytest.approx(prof.series["pe"][-1])
    assert find(ev, "Ratios", "Revenue growth", 1) == "—"  # no prior year
    assert find(ev, "Ratios", "Current ratio", last) == "n/a"  # inputs missing: shown, not zero
    wb = load_workbook(io.BytesIO(build_excel(inp)))
    assert {"Cover", "Statements", "Ratios", "Red flags", "DCF", "Memo"} <= set(wb.sheetnames)


def test_word_report_contents():
    d = docx.Document(io.BytesIO(build_word(report_input())))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Tata Consultancy Services Ltd" in text and "Key metrics" in text
    assert "Valuation (DCF)" in text and "Analyst memo" in text and "Sources and method" in text
    assert d.tables and d.tables[0].rows[0].cells[0].text == "Metric"


def test_report_endpoints(client):
    inp = report_input()
    body = {"company": inp.company, "periods": inp.periods, "units": "crores", "market": inp.market, "dcf": DCF,
            "lines": [{"label": l.label, "statement": l.statement, "key": l.key, "values": l.values, "source": l.source} for l in inp.lines]}
    r = client.post("/api/reports/excel", json=body)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert 'filename="Tata-Consultancy-Services-Ltd-analysis.xlsx"' in r.headers["content-disposition"]
    r = client.post("/api/reports/word", json={**body, "sections": ["statements"]})
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_reports_require_sign_in(anon_client):
    assert anon_client.post("/api/reports/excel", json={}).status_code == 401
