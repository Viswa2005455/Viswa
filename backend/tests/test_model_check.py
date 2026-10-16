"""Excel model checker."""

import io

from openpyxl import Workbook

from app.files.model_check import _constants, _signature, check_workbook


def model_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    ws.append(["Item", "FY23", "FY24", "FY25", "FY26"])
    ws.append(["Units", 10, 12, 14, 16])          # row 2
    ws.append(["Price", 5, 5, 6, 6])              # row 3
    ws.append(["Revenue"] + [f"={c}2*{c}3" for c in "BCDE"])  # row 4
    ws["A5"] = "Gross"
    ws["B5"], ws["C5"], ws["D5"], ws["E5"] = "=B4+B3", "=C4+C3", "=D4+D2", "=E4+E3"  # D5 inconsistent
    ws["A6"] = "Tax"
    ws["B6"] = "=B4*1.18"      # hard-coded rate
    ws["C6"] = "=ROUND(C4,2)"  # structural 2: fine
    ws["D6"] = "=D4/D3-1"      # 1 is allowed
    ws["E6"] = "=E4*12"        # period constant
    ws["A7"] = "Costs"
    ws["B7"], ws["C7"], ws["E7"] = "=B4+B3", "=C4+C3", "=E4+E3"
    ws["D7"] = 500             # typed over a formula row
    ws["A10"], ws["B10"] = "=B10+1", "=A10+1"   # circular pair
    ws["A11"] = "=A11+1"                         # self reference
    ws["A12"] = "=Missing!A1"
    ws["A13"] = '="Hello!"&B2'                   # text with "!" is not a sheet reference
    ws["A14"] = "=#REF!+1"
    ws["A15"] = "#DIV/0!"
    ws["A16"] = "='[Other.xlsx]Sheet1'!A1"
    bs = wb.create_sheet("BS")
    bs.append(["", "FY24", "FY25"])
    bs.append(["Total assets", 100, 200])
    bs.append(["Total liabilities and equity", 100, 210])
    hid = wb.create_sheet("Old")
    hid.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def issues_by(report, category):
    return [i for i in report["issues"] if i["category"] == category]


def test_signature_is_relative():
    assert _signature("=B2*2", 1, 2) == _signature("=C2*2", 1, 3)
    assert _signature("=$B$2*2", 1, 2) == _signature("=$B$2*2", 1, 3)  # absolute refs do not move
    assert _signature("=B2*2", 1, 2) != _signature("=B2*2", 1, 3)
    assert _signature("=$B$2", 1, 2) == _signature("=$B$2", 5, 9)


def test_constants_ignore_structural_arguments():
    assert _constants("=B4*1.18") == [1.18]
    assert _constants("=ROUND(C4,2)") == []
    assert _constants("=VLOOKUP(A1,B1:D9,3,FALSE)") == []
    assert _constants("=B2^2") == []


def test_finds_each_kind_of_issue():
    rep = check_workbook(model_bytes())
    cells = lambda cat: {i["cell"] for i in issues_by(rep, cat)}  # noqa: E731
    assert "B6" in cells("hardcoded_number")
    assert not {"C6", "D6"} & cells("hardcoded_number")
    period = next(i for i in issues_by(rep, "hardcoded_number") if i["cell"] == "E6")
    assert "period conversions" in period["message"]
    assert cells("inconsistent_formula") == {"D5"}
    assert "D7" in cells("value_in_formula_range")
    circ = issues_by(rep, "circular_reference")
    assert len(circ) == 2 and any("2 cell" in i["message"] for i in circ)
    assert cells("missing_sheet") == {"A12"}
    assert "A14" in cells("broken_reference")
    assert "A15" in cells("error_value")
    assert "A16" in cells("external_link")
    bal = issues_by(rep, "balance_check")
    assert len(bal) == 1 and bal[0]["cell"] == "C2 vs C3"
    assert [i["sheet"] for i in issues_by(rep, "hidden_sheet")] == ["Old"]
    assert rep["issues"][0]["severity"] == "high"
    assert rep["formulas_checked"] > 10


def test_clean_model_has_no_high_issues():
    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "FY24", "FY25"])
    ws.append(["Revenue", 100, 120])
    ws.append(["Growth", None, "=C2/B2-1"])
    buf = io.BytesIO()
    wb.save(buf)
    rep = check_workbook(buf.getvalue())
    assert rep["by_severity"]["high"] == 0 and rep["issues"] == []


def test_model_check_endpoint(client):
    up = client.post("/api/uploads", files=[("files", ("model.xlsx", model_bytes()))]).json()["results"][0]["upload"]
    r = client.post(f"/api/analysis/model-check/{up['id']}")
    assert r.status_code == 200, r.text
    assert r.json()["name"] == "model.xlsx" and r.json()["counts"]["circular_reference"] == 2
    csv = client.post("/api/uploads", files=[("files", ("a.csv", b"x,y\n1,2\n"))]).json()["results"][0]["upload"]
    assert client.post(f"/api/analysis/model-check/{csv['id']}").status_code == 422
    assert client.post("/api/analysis/model-check/not-an-id").status_code == 400
    assert client.post(f"/api/analysis/model-check/{'0' * 32}").status_code == 404
