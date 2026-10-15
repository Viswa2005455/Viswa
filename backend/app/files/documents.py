"""PDF, Word, text, structured data, notebook and source-code extraction.

Uploaded content is treated as untrusted data: nothing here executes code. Python files are
parsed with `ast` (which does not run them), XML is parsed with defusedxml, YAML with
safe_load only.
"""

from __future__ import annotations

import ast
import io
import json
import re
import tomllib

import defusedxml.ElementTree as SafeET
import yaml

from .tabular import analyse_grid, decode_text

MAX_PDF_PAGES = 60
MAX_TEXT_PREVIEW = 60_000

_CURRENCY_PAT = {
    "INR": r"₹|\bINR\b|\bRs\.?\s|\brupees?\b",
    "USD": r"\$|\bUSD\b|\bUS dollars?\b",
    "EUR": r"€|\bEUR\b|\beuros?\b",
    "GBP": r"£|\bGBP\b|\bpounds sterling\b",
    "JPY": r"¥|\bJPY\b|\byen\b",
}
_UNIT_PAT = {
    "thousands": r"\bin thousands\b|\b'000s?\b|\(000s?\)",
    "lakhs": r"\blakhs?\b|\blacs?\b",
    "crores": r"\bcrores?\b|\bcr\.?\b",
    "millions": r"\bin millions\b|\bmn\b|\(\s*m\s*\)",
    "billions": r"\bin billions\b|\bbn\b",
}
_PERIOD_PAT = r"\b(?:FY\s?\d{2,4}(?:-\d{2,4})?|Q[1-4]\s?(?:FY)?\s?\d{2,4}|(?:19|20)\d{2})\b"


def detect_financial_context(text: str) -> dict:
    """Detect currencies, units and periods mentioned in text. These are hints for the user to confirm."""
    currencies = {k: len(re.findall(p, text, re.I)) for k, p in _CURRENCY_PAT.items()}
    units = {k: len(re.findall(p, text, re.I)) for k, p in _UNIT_PAT.items()}
    periods = sorted({re.sub(r"\s+", " ", m) for m in re.findall(_PERIOD_PAT, text)})[:30]
    return {
        "currencies": {k: v for k, v in currencies.items() if v},
        "units": {k: v for k, v in units.items() if v},
        "periods": periods,
        "needs_confirmation": True,
    }


def parse_pdf(data: bytes) -> dict:
    import pdfplumber

    pages, sheets, warnings = [], [], []
    all_text = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        total = len(pdf.pages)
        if total > MAX_PDF_PAGES:
            warnings.append(f"Only the first {MAX_PDF_PAGES} of {total} pages were processed.")
        empty_pages = []
        for i, page in enumerate(pdf.pages[:MAX_PDF_PAGES]):
            text = page.extract_text() or ""
            all_text.append(text)
            if not text.strip():
                empty_pages.append(i + 1)
            tables = []
            for t_idx, table in enumerate(page.extract_tables() or []):
                if not table or len(table) < 2:
                    continue
                grid = [[(c.strip() if isinstance(c, str) else c) or None for c in row] for row in table]
                s = analyse_grid(f"Page {i + 1} · Table {t_idx + 1}", grid, excel=False)
                s["source"] = {"page": i + 1, "table": t_idx + 1}
                sheets.append(s)
                tables.append(t_idx + 1)
            pages.append({"page": i + 1, "chars": len(text), "text": text[:4000], "tables": tables})
        if empty_pages:
            warnings.append(f"No extractable text on page(s) {', '.join(map(str, empty_pages[:20]))}. "
                            "They may be scanned images; OCR is not configured on this server.")
    joined = "\n".join(all_text)
    if sheets:
        warnings.append("Numbers extracted from PDF tables can be misaligned. Confirm figures against the source page before relying on them.")
    return {"kind": "document", "page_count": total, "pages": pages, "sheets": sheets,
            "text": joined[:MAX_TEXT_PREVIEW], "context": detect_financial_context(joined), "warnings": warnings}


def parse_docx(data: bytes) -> dict:
    import docx

    d = docx.Document(io.BytesIO(data))
    paragraphs = [p.text for p in d.paragraphs if p.text.strip()]
    sheets = []
    for t_idx, table in enumerate(d.tables):
        grid = [[(cell.text.strip() or None) for cell in row.cells] for row in table.rows]
        if len(grid) >= 2:
            s = analyse_grid(f"Table {t_idx + 1}", grid, excel=False)
            s["source"] = {"table": t_idx + 1}
            sheets.append(s)
    text = "\n".join(paragraphs)
    return {"kind": "document", "paragraph_count": len(paragraphs), "sheets": sheets,
            "text": text[:MAX_TEXT_PREVIEW], "context": detect_financial_context(text), "warnings": []}


def parse_text(data: bytes, label: str) -> dict:
    text, warn = decode_text(data)
    return {"kind": "document", "text": text[:MAX_TEXT_PREVIEW], "line_count": text.count("\n") + 1,
            "context": detect_financial_context(text), "warnings": [warn] if warn else [],
            "truncated": len(text) > MAX_TEXT_PREVIEW, "format": label}


def parse_structured(data: bytes, ext: str, label: str) -> dict:
    text, warn = decode_text(data)
    out = {"kind": "data", "format": label, "text": text[:MAX_TEXT_PREVIEW], "line_count": text.count("\n") + 1,
           "warnings": [warn] if warn else [], "valid": True, "error": None, "structure": None}
    try:
        if ext == ".json":
            obj = json.loads(text)
            out["structure"] = _describe(obj)
            if isinstance(obj, list) and obj and all(isinstance(x, dict) for x in obj[:50]):
                keys = list(dict.fromkeys(k for x in obj[:5000] for k in x))[:200]
                grid = [keys] + [[x.get(k) if not isinstance(x.get(k), (dict, list)) else json.dumps(x.get(k)) for k in keys]
                                 for x in obj[:5000]]
                out["sheets"] = [analyse_grid("Records", grid, excel=False)]
        elif ext == ".xml":
            root = SafeET.fromstring(text.encode())
            out["structure"] = {"root": root.tag, "children": len(list(root))}
        elif ext in (".yml", ".yaml"):
            out["structure"] = _describe(yaml.safe_load(text))
        elif ext == ".toml":
            out["structure"] = _describe(tomllib.loads(text))
    except Exception as e:  # report syntax problems as data, not as a server error
        out["valid"] = False
        out["error"] = f"{type(e).__name__}: {str(e)[:300]}"
    return out


def _describe(obj) -> dict:
    if isinstance(obj, dict):
        return {"type": "object", "keys": list(obj)[:50], "key_count": len(obj)}
    if isinstance(obj, list):
        return {"type": "array", "length": len(obj), "item_type": type(obj[0]).__name__ if obj else None}
    return {"type": type(obj).__name__}


def parse_notebook(data: bytes) -> dict:
    text, _ = decode_text(data)
    try:
        nb = json.loads(text)
    except json.JSONDecodeError as e:
        return {"kind": "code", "language": "Jupyter notebook", "valid": False, "error": str(e), "text": text[:MAX_TEXT_PREVIEW], "warnings": []}
    cells = nb.get("cells", []) if isinstance(nb, dict) else []
    lang = (nb.get("metadata", {}).get("kernelspec", {}) or {}).get("language") or \
           (nb.get("metadata", {}).get("language_info", {}) or {}).get("name") or "python"
    out_cells, syntax_errors = [], []
    for i, c in enumerate(cells[:500]):
        src = "".join(c.get("source", [])) if isinstance(c.get("source"), list) else str(c.get("source", ""))
        out_cells.append({"index": i + 1, "type": c.get("cell_type"), "source": src[:5000],
                          "has_outputs": bool(c.get("outputs"))})
        if c.get("cell_type") == "code" and lang == "python":
            # strip IPython magics/shell lines before checking syntax
            cleaned = "\n".join("" if ln.lstrip().startswith(("%", "!")) else ln for ln in src.splitlines())
            try:
                ast.parse(cleaned)
            except SyntaxError as e:
                syntax_errors.append({"cell": i + 1, "line": e.lineno, "message": e.msg})
    counts = {t: sum(1 for c in cells if c.get("cell_type") == t) for t in ("code", "markdown", "raw")}
    return {"kind": "code", "language": f"Jupyter notebook ({lang})", "valid": True, "cells": out_cells,
            "cell_counts": counts, "syntax_errors": syntax_errors, "warnings": [],
            "note": "Notebook outputs were not executed or re-run."}


def _python_outline(src: str) -> dict:
    try:
        tree = ast.parse(src)
    except SyntaxError as e:
        return {"syntax_ok": False, "syntax_errors": [{"line": e.lineno, "column": e.offset, "message": e.msg}]}
    imports, functions, classes = set(), [], []
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imports.update(a.name.split(".")[0] for a in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            imports.add(node.module.split(".")[0])
        elif isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            functions.append({"name": node.name, "line": node.lineno})
        elif isinstance(node, ast.ClassDef):
            classes.append({"name": node.name, "line": node.lineno})
    risky = sorted(imports & {"os", "subprocess", "socket", "shutil", "requests", "urllib", "ctypes", "pickle"})
    return {"syntax_ok": True, "syntax_errors": [], "imports": sorted(imports), "functions": functions[:200],
            "classes": classes[:100], "risky_imports": risky}


def _sql_outline(src: str) -> dict:
    no_comments = re.sub(r"--[^\n]*|/\*.*?\*/", "", src, flags=re.S)
    statements = [s for s in (x.strip() for x in no_comments.split(";")) if s]
    up = no_comments.upper()
    features = [name for name, pat in {
        "SELECT": r"\bSELECT\b", "JOIN": r"\bJOIN\b", "GROUP BY": r"\bGROUP\s+BY\b", "CTE (WITH)": r"^\s*WITH\b|\bWITH\s+\w+\s+AS\s*\(",
        "Subquery": r"\(\s*SELECT\b", "Window function": r"\bOVER\s*\(", "INSERT/UPDATE/DELETE": r"\b(INSERT|UPDATE|DELETE)\b",
        "DDL (CREATE/ALTER/DROP)": r"\b(CREATE|ALTER|DROP)\b"}.items() if re.search(pat, up, re.M)]
    tables = sorted(set(re.findall(r"\b(?:FROM|JOIN|INTO|UPDATE)\s+([A-Za-z_][\w.]*)", no_comments, re.I)))[:100]
    issues = []
    if no_comments.count("(") != no_comments.count(")"):
        issues.append("Unbalanced parentheses.")
    if no_comments.count("'") % 2:
        issues.append("Unclosed string literal (odd number of single quotes).")
    if re.search(r"\bSELECT\s+\*", up):
        issues.append("SELECT * used; list columns explicitly for stable results.")
    if re.search(r"\b(UPDATE|DELETE)\b(?![^;]*\bWHERE\b)", up):
        issues.append("UPDATE/DELETE without WHERE affects every row.")
    return {"statement_count": len(statements), "features": features, "tables": tables, "issues": issues}


def _java_outline(src: str) -> dict:
    classes = re.findall(r"\b(?:class|interface|enum|record)\s+([A-Z]\w*)", src)
    methods = re.findall(r"(?:public|private|protected|static|\s)+[\w<>\[\],\s]+\s+(\w+)\s*\([^)]*\)\s*(?:throws [\w.,\s]+)?\{", src)
    issues = []
    if src.count("{") != src.count("}"):
        issues.append("Unbalanced braces.")
    return {"classes": classes[:100], "methods": [m for m in methods if m not in ("if", "for", "while", "switch", "catch")][:200],
            "has_main": bool(re.search(r"static\s+void\s+main\s*\(", src)), "issues": issues}


def parse_code(data: bytes, ext: str, label: str) -> dict:
    text, warn = decode_text(data)
    out = {"kind": "code", "language": label, "extension": ext, "text": text[:MAX_TEXT_PREVIEW],
           "line_count": text.count("\n") + 1, "warnings": [warn] if warn else [],
           "note": "Code is never executed on upload. Execution requires explicit approval in the Data & Code Lab."}
    if ext == ".py":
        out["outline"] = _python_outline(text)
    elif ext == ".sql":
        out["outline"] = _sql_outline(text)
    elif ext == ".java":
        out["outline"] = _java_outline(text)
    return out
