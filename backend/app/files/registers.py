"""Read sales / purchase registers (Excel or CSV) into invoice rows for the GST tools.

Column headers are matched by name, ignoring case, spaces and punctuation, so exports from most
accounting packages work without renaming. The GSTR-2B Excel downloaded from the GST portal has
two header rows ("Invoice details" / "Invoice number"); headers spanning two rows are combined.
Files are read in memory and never stored.
"""

from __future__ import annotations

import datetime as dt
import io
import re
from typing import Any

from openpyxl import load_workbook

from .archive import check_zip_container
from .tabular import decode_text, parse_number

MAX_ROWS = 100_000
FIELDS: dict[str, list[str]] = {
    "gstin": ["gstin", "gstinofsupplier", "gstinuinofrecipient", "gstinuin", "customergstin", "recipientgstin", "suppliergstin",
              "partygstin", "buyergstin", "gstno", "gstnumber"],
    "number": ["invoicenumber", "invoiceno", "invno", "documentnumber", "documentno", "docno", "billno", "billnumber",
               "notenumber", "notenumber", "voucherno", "invoicedetailsinvoicenumber", "creditdebitnotedetailsnotenumber"],
    "date": ["invoicedate", "date", "documentdate", "docdate", "billdate", "notedate", "invoicedetailsinvoicedate"],
    "invoice_value": ["invoicevalue", "totalinvoicevalue", "invoicevaluers", "grossvalue", "totalvalue", "invoicedetailsinvoicevalue"],
    "place_of_supply": ["placeofsupply", "pos", "placeofsupplycode", "statecode", "posstatecode"],
    "taxable_value": ["taxablevalue", "taxablevaluers", "taxableamount", "assessablevalue", "taxable"],
    "rate": ["rate", "gstrate", "taxrate", "rateofgst", "ratepercent"],
    "igst": ["igst", "integratedtax", "integratedtaxrs", "igstamount", "taxamountintegratedtax"],
    "cgst": ["cgst", "centraltax", "centraltaxrs", "cgstamount", "taxamountcentraltax"],
    "sgst": ["sgst", "utgst", "statetax", "stateuttax", "stateuttaxrs", "sgstutgst", "sgstamount", "taxamountstateuttax"],
    "cess": ["cess", "cessrs", "cessamount", "taxamountcess"],
    "hsn": ["hsn", "hsncode", "hsnsac", "sac", "hsnsaccode"],
    "doc_type": ["documenttype", "doctype", "type", "notetype", "vouchertype"],
    "category": ["category", "supplycategory", "supplytype"],
}
_ALIASES = {alias: f for f, aliases in FIELDS.items() for alias in aliases}


def _key(v: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(v or "").lower().replace("₹", ""))


def _grids(filename: str, data: bytes) -> list[tuple[str, list[list[Any]]]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        check_zip_container(data, 10_000, 200 * 1024 * 1024)
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            out = []
            for ws in wb.worksheets:
                grid = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= MAX_ROWS + 10:
                        break
                    grid.append(list(row[:60]))
                out.append((ws.title, grid))
            return out
        finally:
            wb.close()
    if name.endswith((".csv", ".txt", ".tsv")):
        import csv

        text, _ = decode_text(data)
        try:
            delim = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
        except csv.Error:
            delim = ","
        rows = []
        for i, r in enumerate(csv.reader(io.StringIO(text), delimiter=delim)):
            if i >= MAX_ROWS + 10:
                break
            rows.append([c if c.strip() else None for c in r[:60]])
        return [("Data", rows)]
    raise ValueError("Upload the register as .xlsx or .csv.")


def _header(grid: list[list[Any]]) -> tuple[int, dict[int, str]] | None:
    """Find the header row (it may sit below a title block) and map columns to fields."""
    best = None
    for r in range(min(len(grid), 15)):
        row = grid[r]
        above = grid[r - 1] if r else []
        mapping: dict[int, str] = {}
        for c, v in enumerate(row):
            k = _key(v)
            parent = _key(above[c]) if c < len(above) else ""
            for cand in (k, parent + k):
                f = _ALIASES.get(cand)
                if f and f not in mapping.values():
                    mapping[c] = f
                    break
        if "number" in mapping.values() and ("taxable_value" in mapping.values() or "invoice_value" in mapping.values()):
            if best is None or len(mapping) > len(best[1]):
                best = (r, mapping)
    return best


def _cell(v: Any, field: str):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if field in ("taxable_value", "rate", "igst", "cgst", "sgst", "cess", "invoice_value"):
        n, _ = parse_number(v)
        return n
    if field == "date":
        if isinstance(v, (dt.date, dt.datetime)):
            return (v.date() if isinstance(v, dt.datetime) else v).isoformat()
        return str(v).strip()
    if field == "place_of_supply":
        s = str(v).strip()
        if isinstance(v, float) and v.is_integer():
            s = str(int(v))
        m = re.match(r"^(\d{1,2})", s)
        return m.group(1).zfill(2) if m else s
    if field == "hsn" and isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_register(filename: str, data: bytes) -> dict:
    """Return {"sheet", "rows": [dict], "columns": {field: header}, "skipped"}; raises ValueError when no table is found."""
    found = None
    for title, grid in _grids(filename, data):
        h = _header(grid)
        if h and (found is None or len(h[1]) > len(found[2])):
            found = (title, grid, h[1], h[0])
    if not found:
        raise ValueError("No invoice table found. The file needs columns such as 'Invoice number' and 'Taxable value' "
                         "(download the template for the expected layout).")
    title, grid, mapping, hdr = found
    rows, skipped = [], 0
    for r in range(hdr + 1, len(grid)):
        vals = grid[r]
        rec = {f: _cell(vals[c] if c < len(vals) else None, f) for c, f in mapping.items()}
        if not rec.get("number"):
            if any(v not in (None, "") for v in rec.values()):
                skipped += 1  # e.g. a totals row
            continue
        if len(rows) >= MAX_ROWS:
            raise ValueError(f"At most {MAX_ROWS:,} rows per file.")
        rec["row"] = r + 1
        rows.append(rec)
    header_row = grid[hdr]
    return {"sheet": title, "rows": rows, "columns": {f: str(header_row[c]) for c, f in mapping.items()}, "skipped": skipped}


SALES_TEMPLATE = ("Invoice number,Invoice date,Customer GSTIN,Place of supply,Document type,Category,HSN,Taxable value,Rate,IGST,CGST,SGST,Cess\n"
                  "INV-001,2026-09-02,27AAPFU0939F1ZV,27,invoice,regular,8471,100000,18,,9000,9000,0\n"
                  "INV-002,2026-09-05,,29,invoice,regular,8471,150000,18,27000,,,0\n"
                  "INV-003,2026-09-07,,27,invoice,regular,8517,20000,18,,1800,1800,0\n"
                  "CN-001,2026-09-20,27AAPFU0939F1ZV,27,credit_note,regular,8471,10000,18,,900,900,0\n")
PURCHASE_TEMPLATE = ("Supplier GSTIN,Invoice number,Invoice date,Taxable value,IGST,CGST,SGST,Cess\n"
                     "27AAPFU0939F1ZV,S/2026/101,2026-09-03,50000,0,4500,4500,0\n")


# ---------------------------------------------------------------- daily price files
_MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
PRICE_FIELDS = {"date": ["date", "tradedate", "timestamp", "day"],
                "close": ["adjclose", "adjustedclose", "close", "closeprice", "closingprice", "ltp", "price", "lasttradedprice"]}


def parse_date(v: Any) -> str | None:
    """ISO, dd-Mon-yyyy (NSE/BSE downloads), dd/mm/yyyy and dd-mm-yyyy. Day-first, as Indian exchanges publish."""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v or "").strip()
    if not s:
        return None
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        y, mo, d = map(int, m.groups())
    else:
        m = re.match(r"^(\d{1,2})[-/ ]([A-Za-z]{3})[A-Za-z]*[-/ ,]+(\d{2,4})$", s)
        if m:
            d, mo, y = int(m.group(1)), _MONTHS.get(m.group(2).lower()), int(m.group(3))
            if mo is None:
                return None
        else:
            m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})$", s)
            if not m:
                return None
            d, mo, y = map(int, m.groups())
        if y < 100:
            y += 2000
    try:
        return dt.date(y, mo, d).isoformat()
    except ValueError:
        return None


def read_prices(filename: str, data: bytes) -> dict:
    """A table of daily prices: a date column and a closing-price column (adjusted close preferred)."""
    for title, grid in _grids(filename, data):
        for r in range(min(len(grid), 15)):
            keys = [_key(v) for v in grid[r]]
            date_col = next((i for i, k in enumerate(keys) if k in PRICE_FIELDS["date"]), None)
            close_col = next((keys.index(a) for a in PRICE_FIELDS["close"] if a in keys), None)
            if date_col is None or close_col is None:
                continue
            dates, closes, bad = [], [], 0
            for row in grid[r + 1:]:
                if date_col >= len(row):
                    continue
                d = parse_date(row[date_col])
                if d is None:
                    bad += 1 if any(c not in (None, "") for c in row) else 0
                    continue
                c = parse_number(row[close_col])[0] if close_col < len(row) else None
                dates.append(d)
                closes.append(c)
            if dates:
                return {"sheet": title, "dates": dates, "closes": closes, "date_column": str(grid[r][date_col]),
                        "price_column": str(grid[r][close_col]), "unreadable_rows": bad}
    raise ValueError("No price table found: the file needs a 'Date' column and a 'Close' (or 'Adj Close') column.")
