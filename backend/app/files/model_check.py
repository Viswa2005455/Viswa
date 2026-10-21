"""Excel model checker: finds common modelling mistakes in an uploaded workbook.

Checks (each issue names the sheet and cell):
* error values (#DIV/0!, #REF!, #N/A, ...) and formulas that contain #REF!;
* links to other workbooks and references to sheets that do not exist;
* circular references;
* numbers typed inside formulas (e.g. =B5*1.18) instead of linking to an input cell;
* a formula that differs from the matching formulas on both sides of it (copy-across mistakes);
* a typed number sitting in a row or column of identical formulas (a formula overwritten by a value);
* balance sheets whose "Total assets" row does not equal "Total liabilities and equity";
* hidden sheets, which can hide inputs or old workings.

Nothing in the workbook is executed: formulas are tokenised by formula_eval's hand-written
tokenizer, and only the cached values Excel saved are read.
"""

from __future__ import annotations

import io
import re
from collections import defaultdict
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .formula_eval import Unsupported, _rc, _split_ref, tokenize

MAX_ROWS = 5000
MAX_COLS = 200
MAX_ISSUES = 500
MAX_RANGE_EDGES = 20_000  # cells expanded from one range when building the dependency graph
ERRORS = {"#DIV/0!", "#REF!", "#N/A", "#VALUE!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
ALLOWED_CONSTANTS = {0.0, 1.0, 2.0, 100.0}
PERIOD_CONSTANTS = {4.0, 12.0, 52.0, 360.0, 365.0}
# Numeric arguments of these functions (after the first) are structural (digits, column numbers, match types).
STRUCTURAL = {"ROUND", "ROUNDUP", "ROUNDDOWN", "MROUND", "VLOOKUP", "HLOOKUP", "XLOOKUP", "INDEX", "MATCH", "XMATCH", "OFFSET",
              "LEFT", "RIGHT", "MID", "TEXT", "CHOOSE", "LARGE", "SMALL", "DATE", "EDATE", "EOMONTH", "WEEKDAY", "FIXED",
              "INDIRECT", "ADDRESS", "SUBTOTAL", "AGGREGATE", "RANK", "PERCENTILE", "QUARTILE", "TRUNC", "SEQUENCE"}
SEVERITY_ORDER = {"high": 0, "medium": 1, "low": 2}
_EXTERNAL = re.compile(r"\[[^\]]+\]")
_STRINGS = re.compile(r'"(?:[^"]|"")*"')
_BS_ASSETS = re.compile(r"^\s*total\s+assets\s*$", re.I)
_BS_LE = re.compile(r"^\s*total\s+(liabilities\s*(and|&)\s*(shareholders'?\s*|stockholders'?\s*)?equity|equity\s*(and|&)\s*liabilities)\s*$", re.I)


def _addr(r: int, c: int) -> str:
    return f"{get_column_letter(c + 1)}{r + 1}"


class _Sheet:
    def __init__(self, title: str, state: str):
        self.title, self.state = title, state
        self.values: dict[tuple[int, int], Any] = {}
        self.formulas: dict[tuple[int, int], str] = {}
        self.truncated = False


def _read(data: bytes) -> list[_Sheet]:
    wb_f = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    wb_v = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    try:
        for ws_f in wb_f.worksheets:
            ws_v = wb_v[ws_f.title]
            sh = _Sheet(ws_f.title, getattr(ws_f, "sheet_state", "visible") or "visible")
            for r, (row_f, row_v) in enumerate(zip(ws_f.iter_rows(max_col=MAX_COLS), ws_v.iter_rows(max_col=MAX_COLS))):
                if r >= MAX_ROWS:
                    sh.truncated = True
                    break
                for c, (cf, cv) in enumerate(zip(row_f, row_v)):
                    fv = getattr(cf, "value", None)
                    if isinstance(fv, str) and fv.startswith("="):
                        sh.formulas[(r, c)] = fv
                    v = getattr(cv, "value", None)
                    if v is not None:
                        sh.values[(r, c)] = v
            if ws_f.max_column and ws_f.max_column > MAX_COLS:
                sh.truncated = True
            sheets.append(sh)
    finally:
        wb_f.close()
        wb_v.close()
    return sheets


# --------------------------------------------------------------------------- formula analysis

def _signature(formula: str, r: int, c: int) -> str | None:
    """Relative (R1C1-style) form of a formula, so =B2*2 in C2 and =C2*2 in D2 compare equal."""
    try:
        tokens = tokenize(formula)
    except Unsupported:
        return None
    out = []
    for kind, text in tokens:
        if kind == "ref":
            parts = []
            for part in text.split(":"):
                sheet, addr = (part.rsplit("!", 1) if "!" in part else ("", part))
                m = re.match(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)$", addr)
                if not m:
                    return None
                ca, col, ra, row = m.groups()
                rr, cc = _rc(f"{col.upper()}{row}")
                cs = f"C{cc}" if ca else f"C[{cc - c}]"
                rs = f"R{rr}" if ra else f"R[{rr - r}]"
                parts.append(f"{sheet.lower()}!{rs}{cs}" if sheet else f"{rs}{cs}")
            out.append(":".join(parts))
        else:
            out.append(text.upper())
    return " ".join(out)


def _constants(formula: str) -> list[float]:
    """Numeric literals typed in a formula, ignoring structural arguments such as ROUND(x, 2)."""
    try:
        tokens = tokenize(formula)
    except Unsupported:
        s = re.sub(r'"(?:[^"]|"")*"', "", formula)
        s = re.sub(r"(?:'(?:[^']|'')+'|[A-Za-z_][\w.]*)!", "", s)
        s = re.sub(r"\$?[A-Za-z]{1,3}\$?\d+|\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}|\$?\d+:\$?\d+|[A-Za-z_][\w.]*", "", s)
        return [float(m) for m in re.findall(r"(?<![\w.])\d+(?:\.\d+)?(?![\w.])", s)]
    found, stack, prev = [], [], None
    for i, (kind, text) in enumerate(tokens):
        if kind == "func":
            stack.append([text.upper(), 0])
            prev = kind
            continue
        if text == "(":
            if prev != "func":
                stack.append([None, 0])
        elif text == ")":
            if stack:
                stack.pop()
        elif text == "," and stack:
            stack[-1][1] += 1
        elif kind == "num":
            top = stack[-1] if stack else None
            structural = top is not None and top[0] in STRUCTURAL and (top[1] >= 1 or top[0] in ("CHOOSE",))
            prev_op = tokens[i - 1][1] if i else ""
            if not structural and not (prev_op == "^" and text in ("2", "0.5")):
                found.append(float(text))
        prev = kind
    return found


def _refs(formula: str, sheet: str) -> list[tuple[str, int, int]] | None:
    """Cells a formula depends on (ranges expanded up to a limit); None if it cannot be tokenised."""
    try:
        tokens = tokenize(formula)
    except Unsupported:
        return None
    out: list[tuple[str, int, int]] = []
    for kind, text in tokens:
        if kind != "ref":
            continue
        sh, addr = _split_ref(text, sheet)
        try:
            if ":" in addr:
                a, b = addr.split(":")
                (r1, c1), (r2, c2) = _rc(a), _rc(b)
                r1, r2 = sorted((r1, r2))
                c1, c2 = sorted((c1, c2))
                if (r2 - r1 + 1) * (c2 - c1 + 1) > MAX_RANGE_EDGES:
                    continue
                out.extend((sh, r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1))
            else:
                r, c = _rc(addr)
                out.append((sh, r, c))
        except Unsupported:
            continue
    return out


def _cycles(sheets: list[_Sheet]) -> list[list[tuple[str, int, int]]]:
    """Strongly connected components with more than one cell (or a self-reference), via iterative Tarjan."""
    lookup = {s.title.lower(): s.title for s in sheets}
    graph: dict[tuple[str, int, int], list[tuple[str, int, int]]] = {}
    formula_cells = {(s.title, r, c) for s in sheets for (r, c) in s.formulas}
    for s in sheets:
        for (r, c), f in s.formulas.items():
            refs = _refs(f, s.title) or []
            graph[(s.title, r, c)] = [(lookup.get(sh.lower(), sh), rr, cc) for sh, rr, cc in refs
                                      if (lookup.get(sh.lower(), sh), rr, cc) in formula_cells]
    index: dict = {}
    low: dict = {}
    on_stack: set = set()
    stack: list = []
    out: list[list[tuple[str, int, int]]] = []
    counter = 0
    for root in graph:
        if root in index:
            continue
        work = [(root, iter(graph[root]))]
        index[root] = low[root] = counter
        counter += 1
        stack.append(root)
        on_stack.add(root)
        while work:
            node, it = work[-1]
            advanced = False
            for nxt in it:
                if nxt not in index:
                    index[nxt] = low[nxt] = counter
                    counter += 1
                    stack.append(nxt)
                    on_stack.add(nxt)
                    work.append((nxt, iter(graph.get(nxt, []))))
                    advanced = True
                    break
                if nxt in on_stack:
                    low[node] = min(low[node], index[nxt])
            if advanced:
                continue
            work.pop()
            if work:
                low[work[-1][0]] = min(low[work[-1][0]], low[node])
            if low[node] == index[node]:
                comp = []
                while True:
                    w = stack.pop()
                    on_stack.discard(w)
                    comp.append(w)
                    if w == node:
                        break
                if len(comp) > 1 or node in graph.get(node, []):
                    out.append(sorted(comp))
    return out


# --------------------------------------------------------------------------- checks

def check_workbook(data: bytes) -> dict:
    sheets = _read(data)
    names = {s.title.lower() for s in sheets}
    issues: list[dict] = []

    def add(category, severity, sheet, cell, message, suggestion="", formula=None):
        issues.append({"category": category, "severity": severity, "sheet": sheet, "cell": cell, "message": message,
                       "suggestion": suggestion, "formula": formula})

    n_formulas = sum(len(s.formulas) for s in sheets)
    for s in sheets:
        if s.state != "visible":
            add("hidden_sheet", "low", s.title, None, f"Sheet '{s.title}' is {s.state}.",
                "Unhide it and check whether it holds inputs or old workings that the model still uses.")
        if s.truncated:
            add("truncated", "low", s.title, None, f"Only the first {MAX_ROWS} rows and {MAX_COLS} columns were checked.")

        # error values and broken references
        for (r, c), v in s.values.items():
            if isinstance(v, str) and v.strip() in ERRORS:
                add("error_value", "high", s.title, _addr(r, c), f"The cell shows {v.strip()}.",
                    {"#DIV/0!": "Guard the division, e.g. =IF(B2=0, \"\", A2/B2), or fix the empty input.",
                     "#REF!": "A referenced cell or sheet was deleted. Re-point the formula.",
                     "#N/A": "A lookup found no match. Check the lookup value and range."}.get(v.strip(), "Trace the inputs with Excel's Trace Precedents."),
                    s.formulas.get((r, c)))
        for (r, c), f in s.formulas.items():
            a = _addr(r, c)
            bare = _STRINGS.sub('""', f)  # text inside quotes is not a reference
            if "#REF!" in bare.upper():
                add("broken_reference", "high", s.title, a, "The formula contains #REF!, so part of it points nowhere.",
                    "Rebuild the reference to the intended cell.", f)
            if _EXTERNAL.search(bare):
                add("external_link", "high", s.title, a, "The formula links to another workbook.",
                    "Copy the needed values into this file (or document the source) so the model works on its own.", f)
            for m in re.finditer(r"(?<![#\w.])(?:'((?:[^']|'')+)'|([A-Za-z_][\w.]*))!", bare):
                ref_sheet = (m.group(1) or m.group(2) or "").replace("''", "'")
                if ref_sheet and not ref_sheet.startswith("[") and ref_sheet.lower() not in names and not _EXTERNAL.search(m.group(0)):
                    add("missing_sheet", "high", s.title, a, f"The formula refers to a sheet called '{ref_sheet}' that is not in this workbook.",
                        "Re-point the formula or restore the sheet.", f)
                    break
            consts = [x for x in _constants(f) if x not in ALLOWED_CONSTANTS]
            if consts:
                period_only = all(x in PERIOD_CONSTANTS for x in consts)
                shown = ", ".join(f"{x:g}" for x in consts[:4])
                add("hardcoded_number", "low", s.title, a,
                    f"Number(s) typed inside the formula: {shown}." + (" These look like period conversions." if period_only else ""),
                    "Put the number in a labelled input cell and refer to it, so it is visible and easy to change.", f)

        # inconsistent formulas and values typed over a formula row/column
        sigs = {(r, c): _signature(f, r, c) for (r, c), f in s.formulas.items()}
        flagged: set[tuple[int, int]] = set()
        for (r, c), sig in sigs.items():
            if sig is None:
                continue
            for dr, dc, direction in ((0, 1, "row"), (1, 0, "column")):
                before, after = sigs.get((r - dr, c - dc)), sigs.get((r + dr, c + dc))
                if before and after and before == after and sig != before and (r, c) not in flagged:
                    flagged.add((r, c))
                    add("inconsistent_formula", "medium", s.title, _addr(r, c),
                        f"This formula differs from the matching formulas on both sides of it in the {direction} "
                        f"({_addr(r - dr, c - dc)} and {_addr(r + dr, c + dc)}).",
                        "Check whether it was edited by mistake; if it is deliberate, add a note explaining why.", s.formulas[(r, c)])
        for (r, c), v in s.values.items():
            if (r, c) in s.formulas or isinstance(v, bool) or not isinstance(v, (int, float)):
                continue
            for dr, dc, direction in ((0, 1, "row"), (1, 0, "column")):
                before, after = sigs.get((r - dr, c - dc)), sigs.get((r + dr, c + dc))
                if before and after and before == after:
                    add("value_in_formula_range", "high", s.title, _addr(r, c),
                        f"A typed number ({v:g}) sits between identical formulas in the {direction}. "
                        "It may have overwritten a formula.",
                        f"Compare with {_addr(r - dr, c - dc)}; restore the formula if the number was pasted over it.")
                    break

        # balance sheet totals
        labels = {r: str(v) for (r, c), v in s.values.items() if c <= 2 and isinstance(v, str)}
        a_rows = [r for r, t in labels.items() if _BS_ASSETS.match(t)]
        le_rows = [r for r, t in labels.items() if _BS_LE.match(t)]
        if a_rows and le_rows:
            ra, rl = a_rows[0], le_rows[0]
            cols = sorted({c for (r, c) in s.values if r in (ra, rl)})
            for c in cols:
                va, vl = s.values.get((ra, c)), s.values.get((rl, c))
                if isinstance(va, (int, float)) and isinstance(vl, (int, float)) and not isinstance(va, bool) and not isinstance(vl, bool):
                    tol = max(1.0, 0.005 * max(abs(va), abs(vl)))
                    if abs(va - vl) > tol:
                        add("balance_check", "high", s.title, f"{_addr(ra, c)} vs {_addr(rl, c)}",
                            f"Total assets ({va:,.2f}) does not equal total liabilities and equity ({vl:,.2f}); "
                            f"difference {va - vl:,.2f}.", "Find the missing or double-counted line.")

    for comp in _cycles(sheets):
        cells = ", ".join(f"{sh}!{_addr(r, c)}" for sh, r, c in comp[:6]) + (" …" if len(comp) > 6 else "")
        sh, r, c = comp[0]
        add("circular_reference", "high", sh, _addr(r, c), f"Circular reference across {len(comp)} cell(s): {cells}.",
            "Break the loop (e.g. use last period's balance for interest) or, if intended, document the iterative calculation.")

    issues.sort(key=lambda i: (SEVERITY_ORDER[i["severity"]], i["category"], i["sheet"], i["cell"] or ""))
    counts: dict[str, int] = defaultdict(int)
    for i in issues:
        counts[i["category"]] += 1
    by_sev = {k: sum(1 for i in issues if i["severity"] == k) for k in SEVERITY_ORDER}
    return {
        "sheets": [{"name": s.title, "state": s.state, "formulas": len(s.formulas), "cells": len(s.values)} for s in sheets],
        "formulas_checked": n_formulas,
        "counts": dict(counts),
        "by_severity": by_sev,
        "issues": issues[:MAX_ISSUES],
        "truncated_issues": max(0, len(issues) - MAX_ISSUES),
        "notes": ["Values are the results Excel last saved. Formulas whose results were never saved cannot be checked for error values.",
                  "Nothing in the file was run: formulas are read as text and analysed."],
    }
