"""Tabular extraction for Excel and CSV/TSV with cell-level source references."""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from collections import Counter
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .formula_eval import FormulaError, Unsupported, WorkbookEvaluator

MAX_ROWS = 5000
MAX_COLS = 200
PREVIEW_ROWS = 60

_NUM_RE = re.compile(r"^\(?-?[\d,]*\.?\d+\)?$")
_CURRENCY = re.compile(r"[₹$€£¥]|\b(?:INR|USD|EUR|GBP|JPY|Rs\.?)\b", re.I)


def parse_number(value: Any) -> tuple[float | None, bool]:
    """Return (number, parsed_from_text). Non-numeric -> (None, False).

    Accepts "1,234.5", "(1,234)" (negative), "-12", "₹ 1,000", "12%" (returned as 12, not 0.12).
    Ambiguous European formats such as "1.234,56" are not guessed.
    """
    if isinstance(value, bool) or value is None:
        return None, False
    if isinstance(value, (int, float)):
        return float(value), False
    if not isinstance(value, str):
        return None, False
    s = _CURRENCY.sub("", value).strip().replace(" ", "").rstrip("%")
    if not s or "," in s and "." in s and s.rfind(",") > s.rfind("."):
        return None, False
    if not _NUM_RE.match(s):
        return None, False
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()").replace(",", "")
    try:
        v = float(s)
    except ValueError:
        return None, False
    return (-v if neg else v), True


def _jsonable(v: Any) -> Any:
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    if isinstance(v, dt.time):
        return v.isoformat()
    if isinstance(v, float) and (v != v or v in (float("inf"), float("-inf"))):
        return None
    if isinstance(v, (int, float, str)) or v is None:
        return v
    return str(v)


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


_PERIOD_HEAD = re.compile(r"^(?:FY|CY)?\s?'?(?:19|20)\d{2}(?:\s?[-/]\s?\d{2,4})?\s?(?:[AEFPB]|PF)?$", re.I)


def _year_value(c: Any) -> int | None:
    if isinstance(c, bool):
        return None
    if isinstance(c, (int, float)) and float(c).is_integer() and 1990 <= c <= 2100:
        return int(c)
    return None


def _period_header_row(grid: list[list[Any]]) -> int | None:
    """A row of period headings ("FY2023A", "2025", 2024 as a number), most cells after the label
    column being periods. Years typed as numbers count only when they run consecutively."""
    for i, row in enumerate(grid[:25]):
        cells = [c for c in row[1:] if not _is_blank(c)]
        if len(cells) < 2:
            continue
        text = [c for c in cells if isinstance(c, str) and _PERIOD_HEAD.match(c.strip())]
        years = [_year_value(c) for c in cells]
        nums = [y for y in years if y is not None]
        consecutive = len(nums) >= 2 and len(set(nums)) == len(nums) and all(abs(a - b) == 1 for a, b in zip(nums, nums[1:]))
        n = len(text) + (len(nums) if consecutive else 0)
        if n >= 2 and n / len(cells) >= 0.6:
            return i
    return None


def _header_row(grid: list[list[Any]]) -> int | None:
    """Index of the header row: a row of period headings if there is one, else the first row that
    looks like a header (>=2 non-empty cells, mostly text)."""
    p = _period_header_row(grid)
    if p is not None:
        return p
    for i, row in enumerate(grid[:25]):
        cells = [c for c in row if not _is_blank(c)]
        if len(cells) >= 2 and sum(isinstance(c, str) and parse_number(c)[0] is None for c in cells) / len(cells) >= 0.6:
            return i
    return None


def analyse_grid(name: str, grid: list[list[Any]], *, excel: bool, formulas: dict[tuple[int, int], str] | None = None,
                 number_formats: dict[int, Counter] | None = None, truncated: bool = False,
                 calculated: set[tuple[int, int]] | None = None) -> dict:
    """Common analysis for one sheet / CSV. `grid` is row-major, 0-based.

    `calculated` holds the cells whose value was computed by formula_eval because the file had
    no saved result for them."""
    formulas = formulas or {}
    calculated = calculated or set()
    n_rows = len(grid)
    n_cols = max((len(r) for r in grid), default=0)
    grid = [list(r) + [None] * (n_cols - len(r)) for r in grid]

    def ref(r: int, c: int) -> str:
        return f"{get_column_letter(c + 1)}{r + 1}" if excel else f"row {r + 1}, col {c + 1}"

    hdr = _header_row(grid)
    data_start = (hdr + 1) if hdr is not None else 0
    headers = [(str(_year_value(grid[hdr][c]) or grid[hdr][c]).strip()) if hdr is not None and not _is_blank(grid[hdr][c])
               else f"Column {get_column_letter(c + 1)}" for c in range(n_cols)]

    columns = []
    for c in range(n_cols):
        vals = [grid[r][c] for r in range(data_start, n_rows)]
        nums = [parse_number(v) for v in vals]
        numeric = sum(1 for v, _ in nums if v is not None)
        text_numbers = sum(1 for v, t in nums if v is not None and t)
        missing = sum(1 for v in vals if _is_blank(v))
        dates = sum(1 for v in vals if isinstance(v, (dt.date, dt.datetime)))
        dtype = "empty" if missing == len(vals) else "date" if dates and dates >= (len(vals) - missing) * 0.6 else \
            "number" if numeric >= (len(vals) - missing) * 0.6 else "text"
        col = {"index": c, "letter": get_column_letter(c + 1) if excel else str(c + 1), "name": headers[c], "dtype": dtype,
               "missing": missing, "numeric": numeric, "numbers_stored_as_text": text_numbers}
        if number_formats and c in number_formats and number_formats[c]:
            col["number_format"] = number_formats[c].most_common(1)[0][0]
        columns.append(col)

    # duplicates among data rows (ignoring fully blank rows)
    seen: Counter = Counter(tuple(_jsonable(v) for v in row) for row in grid[data_start:] if any(not _is_blank(v) for v in row))
    duplicates = sum(n - 1 for n in seen.values() if n > 1)

    # Series candidates (for forecasting / statements): numeric columns and labelled numeric rows.
    series = []
    for col in columns:
        if col["dtype"] == "number" and col["numeric"] >= 3:
            c = col["index"]
            vals, refs, labels = [], [], []
            for r in range(data_start, n_rows):
                v, _ = parse_number(grid[r][c])
                if v is not None:
                    vals.append(v)
                    refs.append(ref(r, c))
                    lab = grid[r][0] if c != 0 else None
                    labels.append(str(_jsonable(lab)) if not _is_blank(lab) else str(r + 1))
            series.append({"id": f"col:{c}", "label": col["name"], "orientation": "column", "values": vals,
                           "refs": refs, "period_labels": labels,
                           "calculated_values": sum(1 for r in range(data_start, n_rows) if (r, c) in calculated)})
    label_col = next((c["index"] for c in columns if c["dtype"] == "text"), None)
    if label_col is not None:
        for r in range(data_start, n_rows):
            label = grid[r][label_col]
            if _is_blank(label) or not isinstance(label, str):
                continue
            vals, refs, plabels, calc = [], [], [], 0
            for c in range(label_col + 1, n_cols):
                v, _ = parse_number(grid[r][c])
                if v is not None:
                    vals.append(v)
                    refs.append(ref(r, c))
                    plabels.append(headers[c])
                    calc += (r, c) in calculated
            if len(vals) >= 2:
                series.append({"id": f"row:{r}", "label": label.strip(), "orientation": "row", "values": vals,
                               "refs": refs, "period_labels": plabels, "calculated_values": calc})
            if len(series) >= 300:
                break

    warnings = []
    if truncated:
        warnings.append(f"Only the first {MAX_ROWS} rows and {MAX_COLS} columns were analysed.")
    if duplicates:
        warnings.append(f"{duplicates} duplicate data row(s) found.")
    if any(c["numbers_stored_as_text"] for c in columns):
        warnings.append("Some numbers are stored as text and were converted; please confirm them.")
    if hdr is None and n_rows:
        warnings.append("No header row detected; generic column names were used.")
    if calculated:
        examples = ", ".join(ref(r, c) for r, c in sorted(calculated)[:3])
        warnings.append(f"{len(calculated)} formula result(s) were not saved in the file, so Finance Intelligence "
                        f"Studio calculated them (e.g. {examples}; shown highlighted). Check a few against Excel "
                        "before relying on them.")
    missing_formula_cache = sum(1 for (r, c) in formulas if r < n_rows and c < n_cols and _is_blank(grid[r][c]))
    if missing_formula_cache:
        warnings.append(f"{missing_formula_cache} formula cell(s) have no result and could not be calculated here, so "
                        "they are left blank (not treated as zero). Open the file in Excel or LibreOffice, save it, "
                        "and upload it again to include them.")

    preview_rows = [[_jsonable(v) for v in row] for row in grid[:PREVIEW_ROWS]]
    return {
        "name": name,
        "n_rows": n_rows,
        "n_cols": n_cols,
        "header_row": (hdr + 1) if hdr is not None else None,
        "columns": columns,
        "preview": {"rows": preview_rows, "first_row_number": 1, "column_letters": [get_column_letter(i + 1) for i in range(n_cols)]},
        "formula_count": len(formulas),
        "formula_samples": [{"cell": ref(r, c), "formula": f} for (r, c), f in list(formulas.items())[:20]],
        "formulas_without_cached_value": missing_formula_cache,
        "formulas_calculated": len(calculated),
        "calculated_cells": [ref(r, c) for r, c in sorted(calculated) if r < PREVIEW_ROWS],
        "duplicate_rows": duplicates,
        "series": series,
        "warnings": warnings,
    }


def parse_xlsx(data: bytes) -> dict:
    wb_f = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    wb_v = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    raw = []  # (title, grid, formulas, number formats, truncated)
    try:
        for ws_f in wb_f.worksheets:
            ws_v = wb_v[ws_f.title]
            grid: list[list[Any]] = []
            formulas: dict[tuple[int, int], str] = {}
            fmts: dict[int, Counter] = {}
            truncated = False
            for r, (row_f, row_v) in enumerate(zip(ws_f.iter_rows(max_col=MAX_COLS), ws_v.iter_rows(max_col=MAX_COLS))):
                if r >= MAX_ROWS:
                    truncated = True
                    break
                vals = []
                for c, (cf, cv) in enumerate(zip(row_f, row_v)):
                    fv = getattr(cf, "value", None)
                    if isinstance(fv, str) and fv.startswith("="):
                        formulas[(r, c)] = fv
                    v = getattr(cv, "value", None)
                    vals.append(v)
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        nf = getattr(cv, "number_format", None)
                        if nf and nf != "General":
                            fmts.setdefault(c, Counter())[nf] += 1
                grid.append(vals)
            if ws_f.max_column and ws_f.max_column > MAX_COLS:
                truncated = True
            raw.append((ws_f.title, grid, formulas, fmts, truncated))
    finally:
        wb_f.close()
        wb_v.close()

    computed = _calculate_missing_results(raw)
    sheets = []
    for title, grid, formulas, fmts, truncated in raw:
        results = computed.get(title, {})
        for (r, c), v in results.items():
            grid[r][c] = v
        # Trim trailing empty rows and columns (read-only mode pads rows to max_col), but never
        # cut off a cell that holds a formula, even when it has no value.
        keep_rows = max((r + 1 for r, _ in formulas), default=0)
        while len(grid) > keep_rows and all(_is_blank(v) for v in grid[-1]):
            grid.pop()
        width = max((max((i + 1 for i, v in enumerate(row) if not _is_blank(v)), default=0) for row in grid), default=0)
        width = max([width] + [c + 1 for r, c in formulas if r < len(grid)])
        grid = [row[:width] for row in grid]
        sheets.append(analyse_grid(title, grid, excel=True, formulas=formulas, number_formats=fmts, truncated=truncated,
                                   calculated={k for k in results if k[0] < len(grid)}))
    return {"kind": "tabular", "sheets": sheets}


def _calculate_missing_results(raw: list) -> dict[str, dict[tuple[int, int], Any]]:
    """Compute formula cells that have no saved result, across all sheets.

    Returns {sheet: {(row, col): value}}. Cells whose formula is unsupported, or whose result is
    an Excel error, are left out, so they stay blank rather than becoming a made-up number."""
    ev = WorkbookEvaluator({t: g for t, g, *_ in raw}, {t: f for t, _, f, *_ in raw})
    out: dict[str, dict[tuple[int, int], Any]] = {}
    for title, grid, formulas, _, _ in raw:
        for (r, c) in sorted(formulas):
            if r >= len(grid) or c >= len(grid[r]) or not _is_blank(grid[r][c]):
                continue
            try:
                v = ev.cell(title, r, c)
            except (Unsupported, FormulaError, RecursionError):
                continue
            if isinstance(v, FormulaError) or _is_blank(v):
                continue
            if isinstance(v, float) and (v != v or v in (float("inf"), float("-inf"))):
                continue
            out.setdefault(title, {})[(r, c)] = v
    return out


def parse_xls(data: bytes) -> dict:
    import xlrd  # optional dependency, listed in pyproject

    book = xlrd.open_workbook(file_contents=data, on_demand=True)
    sheets = []
    for sh in book.sheets():
        rows = min(sh.nrows, MAX_ROWS)
        grid = []
        for r in range(rows):
            row = []
            for c in range(min(sh.ncols, MAX_COLS)):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                else:
                    row.append(cell.value)
            grid.append(row)
        s = analyse_grid(sh.name, grid, excel=True, truncated=sh.nrows > MAX_ROWS or sh.ncols > MAX_COLS)
        s["warnings"].append("Legacy .xls: formulas are not available, only stored values.")
        sheets.append(s)
    return {"kind": "tabular", "sheets": sheets}


def decode_text(data: bytes) -> tuple[str, str | None]:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            if enc == "utf-16" and not (data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff")):
                continue
            return data.decode(enc), None
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "File is not valid UTF-8; decoded as Latin-1. Some characters may be wrong."


def parse_delimited(data: bytes, delimiter: str | None = None) -> dict:
    text, warn = decode_text(data)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    grid, truncated = [], False
    for i, row in enumerate(reader):
        if i >= MAX_ROWS:
            truncated = True
            break
        grid.append([None if not c.strip() else c for c in row[:MAX_COLS]])
    sheet = analyse_grid("Data", grid, excel=False, truncated=truncated)
    if warn:
        sheet["warnings"].append(warn)
    sheet["delimiter"] = {"\t": "tab"}.get(delimiter, delimiter)
    return {"kind": "tabular", "sheets": [sheet]}
