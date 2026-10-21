"""Import a Screener.in "Export to Excel" workbook as a company.

Screener's workbook has a "Data Sheet" holding plain values (no formulas), in ₹ crores, laid out
in blocks that each start with a "Report Date" row:

    PROFIT & LOSS   Report Date | Sales | Raw Material Cost | ... | Net profit | Dividend Amount
    Quarters        Report Date | Sales | Expenses | ...            (quarterly, skipped here)
    BALANCE SHEET   Report Date | Equity Share Capital | Reserves | Borrowings | ... | Cash & Bank | No. of Equity Shares
    CASH FLOW       Report Date | Cash from Operating Activity | ... | Net Cash Flow
    PRICE           one price per annual report date
    DERIVED         Adjusted Equity Shares in Cr

The parser looks for these labels rather than fixed cell positions, keeps the source cell of
every value, and marks every derived line. Nothing is guessed: a label that is not found
simply produces no line.
"""

from __future__ import annotations

import datetime as dt
import io
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Screener label (lower-case) -> (standard key or None, statement, our label)
PL = {
    "sales": ("revenue", "income", "Sales"),
    "raw material cost": (None, "income", "Raw material cost"),
    "change in inventory": (None, "income", "Change in inventory"),
    "power and fuel": (None, "income", "Power and fuel"),
    "other mfr. exp": (None, "income", "Other manufacturing expenses"),
    "employee cost": ("employee_cost", "income", "Employee cost"),
    "selling and admin": ("sga", "income", "Selling and admin"),
    "other expenses": (None, "income", "Other expenses"),
    "other income": ("other_income", "income", "Other income"),
    "depreciation": ("depreciation_amortization", "income", "Depreciation"),
    "interest": ("interest_expense", "income", "Interest"),
    "profit before tax": ("pretax_income", "income", "Profit before tax"),
    "tax": ("income_tax", "income", "Tax"),
    "net profit": ("net_income", "income", "Net profit"),
    "dividend amount": ("dividends", "income", "Dividend amount"),
}
BS = {
    "equity share capital": ("share_capital", "balance", "Equity share capital"),
    "reserves": ("retained_earnings", "balance", "Reserves"),
    "borrowings": ("total_debt", "balance", "Borrowings"),
    "other liabilities": ("other_liabilities", "balance", "Other liabilities"),
    "net block": ("ppe", "balance", "Net block"),
    "capital work in progress": ("cwip", "balance", "Capital work in progress"),
    "investments": ("investments", "balance", "Investments"),
    "other assets": ("other_assets", "balance", "Other assets"),
    "receivables": ("receivables", "balance", "Receivables"),
    "inventory": ("inventory", "balance", "Inventory"),
    "cash & bank": ("cash", "balance", "Cash & bank"),
}
CF = {
    "cash from operating activity": ("operating_cash_flow", "cashflow", "Cash from operating activity"),
    "cash from investing activity": ("investing_cash_flow", "cashflow", "Cash from investing activity"),
    "cash from financing activity": ("financing_cash_flow", "cashflow", "Cash from financing activity"),
    "net cash flow": ("net_change_in_cash", "cashflow", "Net cash flow"),
}


class NotScreener(ValueError):
    pass


def _norm(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().rstrip(":").lower()


def _num(v: Any) -> float | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.replace(",", "").strip()
        try:
            return float(s) if s else None
        except ValueError:
            return None
    return None


def _period(v: Any) -> str | None:
    if isinstance(v, (dt.date, dt.datetime)):
        return v.strftime("%b %Y")
    if isinstance(v, str) and re.search(r"(19|20)\d{2}", v):
        return v.strip()
    return None


def parse_screener(data: bytes) -> dict:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        names = {ws.title.lower(): ws.title for ws in wb.worksheets}
        title = names.get("data sheet")
        if not title:
            raise NotScreener("This does not look like a Screener.in export (no 'Data Sheet' sheet).")
        ws = wb[title]
        grid = [list(r) for r in ws.iter_rows(max_row=400, max_col=40, values_only=True)]
    finally:
        wb.close()

    def label(r: int) -> str:
        return _norm(grid[r][0]) if grid[r] else ""

    def ref(r: int, c: int) -> str:
        return f"Data Sheet!{get_column_letter(c + 1)}{r + 1}"

    company = None
    for r in range(min(len(grid), 10)):
        if label(r) == "company name" and len(grid[r]) > 1 and grid[r][1]:
            company = str(grid[r][1]).strip()

    # find blocks: a "report date" row followed by labelled rows until a blank row or another header
    blocks = []
    for r in range(len(grid)):
        if label(r) == "report date":
            cols = [(c, _period(v)) for c, v in enumerate(grid[r]) if c > 0 and _period(v)]
            if not cols:
                continue
            header = next((label(k) for k in range(r - 1, max(r - 4, -1), -1) if label(k)), "")
            blocks.append({"row": r, "cols": cols, "header": header})
    if not blocks:
        raise NotScreener("No 'Report Date' rows found in the Data Sheet.")

    headers = {"quarters", "balance sheet", "cash flow", "price", "derived", "profit & loss", "meta"}

    def rows_after(b):
        out, blanks = {}, 0
        for r in range(b["row"] + 1, len(grid)):
            lab = label(r)
            if lab == "report date" or lab in headers:
                break
            if not lab:
                blanks += 1
                if blanks >= 2:
                    break
                continue
            blanks = 0
            out.setdefault(lab, r)
        return out

    def values(r, periods, col_of):
        return [_num(grid[r][col_of[p]]) if p in col_of and col_of[p] < len(grid[r]) else None for p in periods]

    def span(r, cols):
        return f"{ref(r, cols[0][0])}–{get_column_letter(cols[-1][0] + 1)}{r + 1}"

    annual: list[str] | None = None
    lines: list[dict] = []
    seen_keys: set[str] = set()
    quarters_skipped = False
    shares_raw = None
    for b in blocks:
        rows = rows_after(b)
        kind = ("quarters" if "quarter" in b["header"] else
                "pl" if "sales" in rows and "net profit" in rows else
                "bs" if "equity share capital" in rows else
                "cf" if "cash from operating activity" in rows else None)
        if kind == "quarters" or (kind == "pl" and annual is not None):
            quarters_skipped = True
            continue
        if kind is None:
            continue
        if annual is None:
            annual = [p for _, p in b["cols"]]
        col_of = {p: c for c, p in b["cols"]}
        for lab, (key, stmt, our) in {"pl": PL, "bs": BS, "cf": CF}[kind].items():
            r = rows.get(lab)
            if r is None:
                continue
            if key in seen_keys:
                key = None
            if key:
                seen_keys.add(key)
            lines.append({"label": our, "statement": stmt, "key": key, "values": values(r, annual, col_of), "source": span(r, b["cols"])})
        if kind == "bs":
            r = rows.get("total")  # the liabilities-side total equals total assets
            if r is not None and "total_assets" not in seen_keys:
                lines.append({"label": "Total (balance sheet)", "statement": "balance", "key": "total_assets",
                              "values": values(r, annual, col_of), "source": span(r, b["cols"])})
                seen_keys.add("total_assets")
            r = rows.get("no. of equity shares")
            if r is not None:
                shares_raw = values(r, annual, col_of)
    if annual is None:
        raise NotScreener("Could not find the annual Profit & Loss block in the Data Sheet.")

    # market data: PRICE row (aligned with the annual columns) and adjusted shares in crores
    first_cols = blocks[0]["cols"]
    market: dict[str, list[float | None]] = {}
    for r in range(len(grid)):
        lab = label(r)
        if lab == "price":
            vals = [_num(grid[r][c]) if c < len(grid[r]) else None for c, _ in first_cols]
            if any(v is not None for v in vals):
                market["share_price"] = vals
        if lab.startswith("adjusted equity shares in cr"):
            vals = [_num(grid[r][c]) if c < len(grid[r]) else None for c, _ in first_cols]
            if any(v is not None for v in vals):
                market["shares_outstanding"] = vals
    notes = ["Screener figures are in ₹ crores. Shares outstanding are in crores, so price × shares gives market cap in ₹ crores."]
    if "shares_outstanding" not in market and shares_raw and any(v for v in shares_raw if v):
        market["shares_outstanding"] = [v / 1e7 if v is not None else None for v in shares_raw]
        notes.append("Shares outstanding were converted from 'No. of Equity Shares' to crores.")

    # derived lines (always labelled)
    def series(key):
        return next((l["values"] for l in lines if l["key"] == key), None)

    def series_label(label):
        return next((l["values"] for l in lines if l["label"] == label), None)

    derived = []
    n = len(annual)
    pbt, interest, sc, res = series("pretax_income"), series("interest_expense"), series("share_capital"), series("retained_earnings")
    if pbt and interest:
        derived.append({"label": "EBIT (derived: profit before tax + interest)", "statement": "income", "key": "operating_income",
                        "values": [p + i if p is not None and i is not None else None for p, i in zip(pbt, interest)], "source": "derived"})
    if sc and res:
        derived.append({"label": "Total equity (derived: share capital + reserves)", "statement": "balance", "key": "total_equity",
                        "values": [a + b if a is not None and b is not None else None for a, b in zip(sc, res)], "source": "derived"})
    debt, other_l = series("total_debt"), series("other_liabilities")
    if debt and other_l:
        derived.append({"label": "Total liabilities (derived: borrowings + other liabilities)", "statement": "balance", "key": "total_liabilities",
                        "values": [a + b if a is not None and b is not None else None for a, b in zip(debt, other_l)], "source": "derived"})
    lines.extend(derived)

    # optional approximations, applied only if the user opts in on import
    approx = {}
    oa = series("other_assets")
    if oa:
        approx["current_assets"] = {"label": "Current assets (approx.: Screener 'Other assets')", "statement": "balance",
                                    "key": "current_assets", "values": oa, "source": "approximation"}
    if other_l:
        approx["current_liabilities"] = {"label": "Current liabilities (approx.: Screener 'Other liabilities')", "statement": "balance",
                                         "key": "current_liabilities", "values": other_l, "source": "approximation"}
    rm, inv, pf, mfr = series_label("Raw material cost"), series_label("Change in inventory"), series_label("Power and fuel"), series_label("Other manufacturing expenses")
    if rm is not None:
        parts = [x for x in (rm, pf, mfr) if x is not None]
        cogs = [None if any(x[i] is None for x in parts) or (inv is not None and inv[i] is None)
                else sum(x[i] for x in parts) - (inv[i] if inv is not None else 0.0) for i in range(n)]
        approx["cogs"] = {"label": "Cost of goods sold (approx.: raw material + power & fuel + other manufacturing − change in inventory)",
                          "statement": "income", "key": "cogs", "values": cogs, "source": "approximation"}
    if quarters_skipped:
        notes.append("Quarterly results were found and skipped; only annual figures are imported.")
    notes.append("Screener does not report cost of goods sold, capital expenditure, current assets or current liabilities "
                 "separately. Tests that need them show 'Needs data' unless you add them or accept the labelled approximations.")
    return {
        "company": company,
        "periods": annual,
        "units": "crores",
        "lines": lines,
        "market": {k: v for k, v in market.items() if len(v) == n},
        "approximations": approx,
        "notes": notes,
    }
