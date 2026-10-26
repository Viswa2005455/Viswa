"""Word draft of the tax audit report (Form 3CB wording + Form 3CD clause table) for the auditor."""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt, RGBColor

from ..finance.taxaudit import _inr


def build_docx(r: dict) -> bytes:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    warn = doc.add_paragraph()
    run = warn.add_run("DRAFT prepared by Finance Intelligence Studio for the auditor's review. Not a signed audit report. "
                       "PAN and GSTIN are masked: enter the full numbers in the e-filing utility.")
    run.bold = True
    run.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)

    app = r["applicability"]
    doc.add_heading("Is a tax audit required? (section 44AB, AY 2026-27)", level=1)
    verdict = {True: f"Yes — {', '.join(app['clauses'])}. Report: {app['forms']}.", False: "No, on the figures given.",
               None: "Cannot decide yet — some figures are missing."}[app["required"]]
    doc.add_paragraph(verdict)
    for t in app["tests"]:
        res = {True: "Audit required", False: "Not required", None: "Missing figures"}[t["result"]]
        doc.add_paragraph(f"{t['clause']}: {t['working']} → {res}", style="List Bullet")
    doc.add_paragraph(r["due"]["message"])
    if app.get("penalty_271b"):
        doc.add_paragraph(f"Maximum penalty under s.271B if not done: {_inr(app['penalty_271b'])}. {app['penalty_note']}")

    cb = r["report_3cb"]
    doc.add_heading(cb["title"], level=1)
    doc.add_paragraph().add_run(cb["note"]).italic = True
    for p in cb["paragraphs"]:
        doc.add_paragraph(p)
    for f in cb["auditor_fields"]:
        doc.add_paragraph(f"{f}: ____________________")
    doc.add_paragraph(cb["auditor_note"])

    doc.add_heading("FORM NO. 3CD — Statement of particulars (draft)", level=1)
    labels = r["status_labels"]
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    for i, h in enumerate(("Clause", "Particulars", "Draft answer", "Status / source")):
        tbl.rows[0].cells[i].text = h
    for c in r["clauses"]:
        cells = tbl.add_row().cells
        cells[0].text = c["no"]
        cells[1].text = c["title"]
        ans = c["answer"] or ""
        if c["rows"]:
            ans += ("\n" if ans else "") + "\n".join(f"{x['metric']}: {x['substituted']} = {x['result']}" for x in c["rows"][:20])
        if c["status"] == "needs_input":
            ans = ans or f"[Needs: {c['needed']}]"
        if c.get("note"):
            ans += f"\nNote: {c['note']}"
        cells[2].text = ans
        srcs = sorted({x["source"] for x in c["rows"] if x.get("source")} | {e.get("file", "") for e in c["evidence"]} - {""})
        cells[3].text = labels.get(c["status"], c["status"]) + (f"\n{'; '.join(srcs)[:300]}" if srcs else "")
    doc.add_heading("Notes", level=2)
    for n in r["notes"]:
        doc.add_paragraph(n, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
