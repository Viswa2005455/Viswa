"""Company analysis pack: Excel workbook with live formulas, and a Word report.

The Excel ratios are formulas that point at the statement cells, and the DCF sheet is fully
formula-driven, so a reviewer can trace and change every figure in Excel. Figures that need
the app's engine (red-flag scores) are written as values and labelled as such.
"""

from __future__ import annotations

import datetime as dt
import io
from dataclasses import dataclass, field

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..finance import forensic, profile
from ..finance.dcf import DCFInputs, run_dcf

NAVY = "0B1B33"
TEAL = "0F766E"
GREY = "F1F5F9"
THIN = Side(style="thin", color="CBD5E1")
PCT = "0.0%"
NUM = "#,##0.00;(#,##0.00);-"
NUM0 = "#,##0;(#,##0);-"
DISCLAIMER = ("Prepared with Finance Intelligence Studio from the data listed under Sources. Calculations follow the formulas "
              "shown; derived and approximated lines are labelled. Screening scores indicate where to investigate and are not "
              "evidence of misstatement or distress. This pack is analysis, not investment advice.")


@dataclass
class Line:
    label: str
    statement: str
    key: str | None
    values: list[float | None]
    source: str | None = None


@dataclass
class ReportInput:
    company: str
    periods: list[str]
    lines: list[Line]
    units: str = "units"
    currency: str = "INR"
    market: dict[str, list[float | None]] = field(default_factory=dict)
    dcf: dict | None = None
    memo: str | None = None
    prepared_by: str = ""
    sections: list[str] = field(default_factory=lambda: ["statements", "ratios", "red_flags", "valuation", "memo"])


def _items(inp: ReportInput) -> dict[str, list[float | None]]:
    items = {l.key: l.values for l in inp.lines if l.key}
    for k, v in inp.market.items():
        if len(v) == len(inp.periods):
            items[k] = v
    return items


# --------------------------------------------------------------------------- Excel

def _hdr(ws, row: int, values: list, fill=NAVY):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center" if c > 1 else "left", vertical="center", wrap_text=True)


def build_excel(inp: ReportInput) -> bytes:
    wb = Workbook()
    n = len(inp.periods)
    unit_label = f"{inp.currency} {inp.units}" if inp.units != "units" else inp.currency

    # Cover
    ws = wb.active
    ws.title = "Cover"
    ws["A1"] = f"{inp.company} — financial analysis pack"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    rows = [("Prepared by", inp.prepared_by or "—"), ("Date", dt.date.today().isoformat()),
            ("Periods", f"{inp.periods[0]} – {inp.periods[-1]}"), ("Amounts", unit_label)]
    for i, (k, v) in enumerate(rows, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    ws.cell(row=8, column=1, value="Contents").font = Font(bold=True)
    contents = {"statements": "Statements — source figures with cell references", "ratios": "Ratios — live formulas linked to Statements",
                "red_flags": "Red flags — Beneish, Altman, Piotroski (engine values)", "valuation": "DCF — live formula model",
                "memo": "Memo — analyst commentary"}
    r = 9
    for s in inp.sections:
        if s in contents and (s != "valuation" or inp.dcf) and (s != "memo" or inp.memo):
            ws.cell(row=r, column=1, value=f"• {contents[s]}")
            r += 1
    ws.cell(row=r + 1, column=1, value=DISCLAIMER).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 4, end_column=6)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40

    # Statements
    st = wb.create_sheet("Statements")
    _hdr(st, 1, ["Line item", "Standard item", *inp.periods, "Source"])
    row_of: dict[str, int] = {}
    r = 2
    for stmt, title in (("income", "INCOME STATEMENT"), ("balance", "BALANCE SHEET"), ("cashflow", "CASH FLOW")):
        lines = [l for l in inp.lines if l.statement == stmt]
        if not lines:
            continue
        st.cell(row=r, column=1, value=title).font = Font(bold=True, color=TEAL)
        r += 1
        for l in lines:
            st.cell(row=r, column=1, value=l.label)
            st.cell(row=r, column=2, value=l.key or "")
            for j, v in enumerate(l.values):
                c = st.cell(row=r, column=3 + j, value=v)
                c.number_format = NUM
            st.cell(row=r, column=3 + n, value=l.source or "")
            if l.source in ("derived", "approximation"):
                st.cell(row=r, column=1).font = Font(italic=True, color="64748B")
            if l.key and l.key not in row_of:
                row_of[l.key] = r
            r += 1
        r += 1
    for k, vals in inp.market.items():
        if len(vals) != n:
            continue
        st.cell(row=r, column=1, value={"share_price": "Share price", "shares_outstanding": "Shares outstanding"}.get(k, k))
        st.cell(row=r, column=2, value=k)
        for j, v in enumerate(vals):
            st.cell(row=r, column=3 + j, value=v).number_format = NUM
        st.cell(row=r, column=3 + n, value="market data")
        row_of[k] = r
        r += 1
    st.column_dimensions["A"].width = 44
    st.column_dimensions["B"].width = 22
    for j in range(n):
        st.column_dimensions[get_column_letter(3 + j)].width = 14
    st.column_dimensions[get_column_letter(3 + n)].width = 26
    st.freeze_panes = "C2"

    def ref(key: str, j: int) -> str | None:
        rr = row_of.get(key)
        return f"Statements!{get_column_letter(3 + j)}{rr}" if rr else None

    # Ratios with live formulas
    if "ratios" in inp.sections:
        ra = wb.create_sheet("Ratios")
        _hdr(ra, 1, ["Ratio", *inp.periods, "Formula"])
        defs = [
            ("Revenue growth", PCT, ["revenue"], lambda j: f"={ref('revenue', j)}/{ref('revenue', j - 1)}-1" if j > 0 else None, "Revenue t / Revenue t−1 − 1"),
            ("EBIT margin", PCT, ["operating_income", "revenue"], lambda j: f"={ref('operating_income', j)}/{ref('revenue', j)}", "EBIT / Revenue"),
            ("EBITDA margin", PCT, ["operating_income", "depreciation_amortization", "revenue"],
             lambda j: f"=({ref('operating_income', j)}+{ref('depreciation_amortization', j)})/{ref('revenue', j)}", "(EBIT + D&A) / Revenue"),
            ("Net margin", PCT, ["net_income", "revenue"], lambda j: f"={ref('net_income', j)}/{ref('revenue', j)}", "Net income / Revenue"),
            ("Return on equity (avg)", PCT, ["net_income", "total_equity"],
             lambda j: f"={ref('net_income', j)}/AVERAGE({ref('total_equity', j)},{ref('total_equity', j - 1)})" if j > 0 else None, "Net income / average equity"),
            ("Return on assets (avg)", PCT, ["net_income", "total_assets"],
             lambda j: f"={ref('net_income', j)}/AVERAGE({ref('total_assets', j)},{ref('total_assets', j - 1)})" if j > 0 else None, "Net income / average total assets"),
            ("Asset turnover (avg)", "0.00x", ["revenue", "total_assets"],
             lambda j: f"={ref('revenue', j)}/AVERAGE({ref('total_assets', j)},{ref('total_assets', j - 1)})" if j > 0 else None, "Revenue / average total assets"),
            ("Debt / equity", "0.00x", ["total_debt", "total_equity"], lambda j: f"={ref('total_debt', j)}/{ref('total_equity', j)}", "Total debt / Total equity"),
            ("Interest cover", "0.0x", ["operating_income", "interest_expense"], lambda j: f"={ref('operating_income', j)}/{ref('interest_expense', j)}", "EBIT / Interest"),
            ("Current ratio", "0.00x", ["current_assets", "current_liabilities"], lambda j: f"={ref('current_assets', j)}/{ref('current_liabilities', j)}", "Current assets / Current liabilities"),
            ("Receivable days", "0", ["receivables", "revenue"], lambda j: f"={ref('receivables', j)}/{ref('revenue', j)}*365", "Receivables / Revenue × 365"),
            ("Cash conversion", "0.00x", ["operating_cash_flow", "net_income"], lambda j: f"={ref('operating_cash_flow', j)}/{ref('net_income', j)}", "Operating cash flow / Net income"),
            ("Dividend payout", PCT, ["dividends", "net_income"], lambda j: f"={ref('dividends', j)}/{ref('net_income', j)}", "Dividends / Net income"),
            ("P/E", "0.0x", ["share_price", "shares_outstanding", "net_income"],
             lambda j: f"={ref('share_price', j)}*{ref('shares_outstanding', j)}/{ref('net_income', j)}", "Price × shares / Net income"),
        ]
        r = 2
        for name, fmt, needs, f, text in defs:
            missing = [k for k in needs if k not in row_of]
            ra.cell(row=r, column=1, value=name)
            for j in range(n):
                cell = ra.cell(row=r, column=2 + j)
                if missing:
                    cell.value = "n/a"
                else:
                    formula = f(j)
                    cell.value = f"=IFERROR({formula[1:]},\"n/a\")" if formula else "—"
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal="right")
            ra.cell(row=r, column=2 + n, value=text if not missing else f"{text} — needs: {', '.join(missing)}")
            r += 1
        ra.column_dimensions["A"].width = 28
        for j in range(n):
            ra.column_dimensions[get_column_letter(2 + j)].width = 13
        ra.column_dimensions[get_column_letter(2 + n)].width = 48
        ra.freeze_panes = "B2"
        ra.cell(row=r + 1, column=1, value="All ratios are live formulas on the Statements sheet; 'n/a' means an input is missing or a denominator is zero.").font = Font(italic=True, color="64748B")

    # Red flags (engine values)
    if "red_flags" in inp.sections and n >= 2:
        rf = wb.create_sheet("Red flags")
        rep = forensic.analyse(inp.periods, _items(inp))
        rf["A1"] = f"Screening scores — {rep.prior} → {rep.current}"
        rf["A1"].font = Font(bold=True, size=13, color=NAVY)
        _hdr(rf, 3, ["Score", "Model", "Value", "Verdict", "Threshold", "Missing inputs"])
        r = 4
        for s in [rep.beneish, *rep.altman]:
            rf.cell(row=r, column=1, value=s.name)
            rf.cell(row=r, column=2, value=s.model)
            rf.cell(row=r, column=3, value=s.value).number_format = "0.00"
            rf.cell(row=r, column=4, value=s.verdict_text)
            rf.cell(row=r, column=5, value=s.thresholds)
            rf.cell(row=r, column=6, value=", ".join(s.missing))
            r += 1
        rf.cell(row=r, column=1, value="Piotroski F-score")
        rf.cell(row=r, column=3, value=f"{rep.piotroski.score} / {rep.piotroski.tested}")
        rf.cell(row=r, column=4, value=rep.piotroski.verdict_text)
        r += 2
        _hdr(rf, r, ["Warning", "Severity", "Period", "Detail", "", ""])
        for f in rep.flags:
            r += 1
            rf.cell(row=r, column=1, value=f.title)
            rf.cell(row=r, column=2, value=f.severity)
            rf.cell(row=r, column=3, value=f.period or "")
            rf.cell(row=r, column=4, value=f.detail)
        rf.cell(row=r + 2, column=1, value="Values computed by the Finance Intelligence Studio engine; component workings are shown in the app.").font = Font(italic=True, color="64748B")
        for col, w in zip("ABCDEF", (30, 26, 10, 60, 40, 40)):
            rf.column_dimensions[col].width = w

    # DCF with live formulas
    if "valuation" in inp.sections and inp.dcf:
        _dcf_sheet(wb, inp.dcf)

    if "memo" in inp.sections and inp.memo:
        me = wb.create_sheet("Memo")
        me["A1"] = "Analyst memo"
        me["A1"].font = Font(bold=True, size=13, color=NAVY)
        for i, para in enumerate([p for p in inp.memo.split("\n")], 3):
            me.cell(row=i, column=1, value=para).alignment = Alignment(wrap_text=True, vertical="top")
        me.column_dimensions["A"].width = 120

    wb.calculation.fullCalcOnLoad = True  # Excel / LibreOffice recalculate every formula when the file opens
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _dcf_sheet(wb: Workbook, m: dict):
    inp = DCFInputs(**m)
    run_dcf(inp)  # validates inputs; the sheet recomputes everything with formulas
    ws = wb.create_sheet("DCF")
    ws["A1"] = "Discounted cash flow (FCFF) — live model"
    ws["A1"].font = Font(bold=True, size=13, color=NAVY)
    a = [("Base-year revenue", inp.base_revenue, NUM), ("Tax rate", inp.tax_rate, PCT), ("D&A % revenue", inp.da_pct_revenue, PCT),
         ("CapEx % revenue", inp.capex_pct_revenue, PCT), ("NWC % revenue", inp.nwc_pct_revenue, PCT), ("WACC", inp.wacc, PCT),
         ("Terminal growth", inp.terminal_growth, PCT), ("Debt", inp.debt, NUM), ("Cash", inp.cash, NUM),
         ("Shares outstanding", inp.shares_outstanding, NUM), ("Exit multiple (if used)", inp.exit_multiple, "0.0x"),
         ("Mid-year convention (1 = yes)", 1 if inp.mid_year else 0, "0")]
    ws["A3"] = "Assumptions (edit the blue cells)"
    ws["A3"].font = Font(bold=True)
    cell_of = {}
    for i, (k, v, f) in enumerate(a, 4):
        ws.cell(row=i, column=1, value=k)
        c = ws.cell(row=i, column=2, value=v)
        c.number_format = f
        c.font = Font(color="1D4ED8")
        cell_of[k] = f"$B${i}"
    n = len(inp.revenue_growth)
    margins = inp.ebit_margin if len(inp.ebit_margin) == n else inp.ebit_margin * n
    top = 4 + len(a) + 1
    _hdr(ws, top, ["", *[f"Year {i}" for i in range(1, n + 1)]])
    labels = ["Revenue growth", "EBIT margin", "Revenue", "EBIT", "Taxes on EBIT", "NOPAT", "D&A", "CapEx", "NWC", "Change in NWC",
              "FCFF", "Discount period", "Discount factor", "PV of FCFF"]
    R = {lab: top + 1 + i for i, lab in enumerate(labels)}
    for lab, rr in R.items():
        ws.cell(row=rr, column=1, value=lab)
    for j in range(n):
        col = get_column_letter(2 + j)
        prev = get_column_letter(1 + j)
        g = ws.cell(row=R["Revenue growth"], column=2 + j, value=inp.revenue_growth[j])
        g.number_format, g.font = PCT, Font(color="1D4ED8")
        mg = ws.cell(row=R["EBIT margin"], column=2 + j, value=margins[j])
        mg.number_format, mg.font = PCT, Font(color="1D4ED8")
        prev_rev = cell_of["Base-year revenue"] if j == 0 else f"{prev}{R['Revenue']}"
        prev_nwc = f"{cell_of['Base-year revenue']}*{cell_of['NWC % revenue']}" if j == 0 else f"{prev}{R['NWC']}"
        f = {
            "Revenue": f"={prev_rev}*(1+{col}{R['Revenue growth']})",
            "EBIT": f"={col}{R['Revenue']}*{col}{R['EBIT margin']}",
            "Taxes on EBIT": f"=MAX({col}{R['EBIT']},0)*{cell_of['Tax rate']}",
            "NOPAT": f"={col}{R['EBIT']}-{col}{R['Taxes on EBIT']}",
            "D&A": f"={col}{R['Revenue']}*{cell_of['D&A % revenue']}",
            "CapEx": f"={col}{R['Revenue']}*{cell_of['CapEx % revenue']}",
            "NWC": f"={col}{R['Revenue']}*{cell_of['NWC % revenue']}",
            "Change in NWC": f"={col}{R['NWC']}-{prev_nwc}",
            "FCFF": f"={col}{R['NOPAT']}+{col}{R['D&A']}-{col}{R['CapEx']}-{col}{R['Change in NWC']}",
            "Discount period": f"={j + 1}-0.5*{cell_of['Mid-year convention (1 = yes)']}",
            "Discount factor": f"=1/(1+{cell_of['WACC']})^{col}{R['Discount period']}",
            "PV of FCFF": f"={col}{R['FCFF']}*{col}{R['Discount factor']}",
        }
        for lab, formula in f.items():
            c = ws.cell(row=R[lab], column=2 + j, value=formula)
            c.number_format = "0.0000" if lab == "Discount factor" else "0.0" if lab == "Discount period" else NUM
    last = get_column_letter(1 + n)
    first = "B"
    out = R["PV of FCFF"] + 2
    tv_formula = (f"={last}{R['FCFF']}*(1+{cell_of['Terminal growth']})/({cell_of['WACC']}-{cell_of['Terminal growth']})"
                  if inp.terminal_method == "gordon" else f"=({last}{R['EBIT']}+{last}{R['D&A']})*{cell_of['Exit multiple (if used)']}")
    results = [
        ("Sum of PV of FCFF", f"=SUM({first}{R['PV of FCFF']}:{last}{R['PV of FCFF']})"),
        ("Terminal value" + (" (Gordon)" if inp.terminal_method == "gordon" else " (exit multiple)"), tv_formula),
        ("PV of terminal value", f"=B{out + 1}/(1+{cell_of['WACC']})^{n}"),
        ("Enterprise value", f"=B{out}+B{out + 2}"),
        ("Less: debt", f"=-{cell_of['Debt']}"),
        ("Add: cash", f"={cell_of['Cash']}"),
        ("Equity value", f"=B{out + 3}+B{out + 4}+B{out + 5}"),
        ("Value per share", f"=IFERROR(B{out + 6}/{cell_of['Shares outstanding']},\"n/a\")"),
        ("Terminal value share of EV", f"=B{out + 2}/B{out + 3}"),
    ]
    for i, (k, f) in enumerate(results):
        ws.cell(row=out + i, column=1, value=k).font = Font(bold=k in ("Enterprise value", "Equity value", "Value per share"))
        c = ws.cell(row=out + i, column=2, value=f)
        c.number_format = PCT if "share of EV" in k else NUM
    ws.column_dimensions["A"].width = 32
    for j in range(n + 1):
        ws.column_dimensions[get_column_letter(2 + j)].width = 14


# --------------------------------------------------------------------------- Word

def _fmt(v, kind="num"):
    if v is None:
        return "—"
    if kind == "pct":
        return f"{v * 100:.1f}%"
    if kind == "x":
        return f"{v:.2f}×"
    return f"{v:,.0f}" if abs(v) >= 100 else f"{v:,.2f}"


def build_word(inp: ReportInput) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    t = doc.add_heading(f"{inp.company}", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph(f"Financial analysis · {inp.periods[0]} – {inp.periods[-1]} · amounts in {inp.currency} {inp.units}"
                          + (f" · prepared by {inp.prepared_by}" if inp.prepared_by else "") + f" · {dt.date.today().isoformat()}")
    p.runs[0].font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    items = _items(inp)
    prof = profile.build(inp.periods, items)
    last = len(inp.periods) - 1

    doc.add_heading("Key metrics", level=1)
    keys = [("revenue", "num"), ("net_income", "num"), ("revenue_growth", "pct"), ("ebitda_margin", "pct"), ("net_margin", "pct"),
            ("roe", "pct"), ("roce", "pct"), ("debt_equity", "x"), ("interest_cover", "x"), ("cash_conversion", "x"), ("pe", "x")]
    show = [k for k, _ in keys if any(v is not None for v in prof.series[k])]
    cols = inp.periods[-5:]
    tbl = doc.add_table(rows=1, cols=1 + len(cols))
    tbl.style = "Light Grid Accent 1"
    tbl.rows[0].cells[0].text = "Metric"
    for i, c in enumerate(cols, 1):
        tbl.rows[0].cells[i].text = c
    offset = len(inp.periods) - len(cols)
    for k, kind in keys:
        if k not in show:
            continue
        row = tbl.add_row().cells
        row[0].text = prof.metrics[k]["label"]
        for i in range(len(cols)):
            row[1 + i].text = _fmt(prof.series[k][offset + i], kind)
    if prof.highlights:
        doc.add_heading("Highlights", level=2)
        for h in prof.highlights:
            doc.add_paragraph(h, style="List Bullet")

    if "red_flags" in inp.sections and len(inp.periods) >= 2:
        rep = forensic.analyse(inp.periods, items)
        doc.add_heading("Earnings quality and distress screens", level=1)
        for s in [rep.beneish, *rep.altman]:
            val = f"{s.value:.2f}" if s.value is not None else "not calculated"
            doc.add_paragraph(f"{s.name} ({s.model}): {val}. {s.verdict_text}", style="List Bullet")
        doc.add_paragraph(f"Piotroski F-score: {rep.piotroski.verdict_text}", style="List Bullet")
        if rep.flags:
            doc.add_heading("Warning signs", level=2)
            for f in rep.flags:
                doc.add_paragraph(f"[{f.severity.upper()}] {f.title}{' — ' + f.period if f.period else ''}: {f.detail}", style="List Bullet")

    if "valuation" in inp.sections and inp.dcf:
        r = run_dcf(DCFInputs(**inp.dcf))
        doc.add_heading("Valuation (DCF)", level=1)
        doc.add_paragraph(f"Enterprise value {_fmt(r.enterprise_value)}; equity value {_fmt(r.equity_value)}"
                          + (f"; value per share {r.value_per_share:,.2f}" if r.value_per_share is not None else "")
                          + f". Terminal value is {r.terminal_value_share_of_ev * 100:.0f}% of enterprise value. "
                          f"WACC {inp.dcf['wacc'] * 100:.2f}%, terminal growth {inp.dcf.get('terminal_growth', 0) * 100:.2f}%.")
        for w in r.warnings:
            doc.add_paragraph(w, style="List Bullet")

    if "memo" in inp.sections and inp.memo:
        doc.add_heading("Analyst memo", level=1)
        for para in inp.memo.split("\n"):
            s = para.strip()
            if not s:
                continue
            if s.startswith("#"):
                doc.add_heading(s.lstrip("#").strip(), level=2)
            elif s[:2] in ("- ", "* ", "• "):
                doc.add_paragraph(s[2:], style="List Bullet")
            else:
                doc.add_paragraph(s.replace("**", ""))

    doc.add_heading("Sources and method", level=1)
    srcs = sorted({l.source for l in inp.lines if l.source and l.source not in ("derived", "approximation")})
    doc.add_paragraph("Source cells: " + ("; ".join(srcs[:12]) + (" …" if len(srcs) > 12 else "") if srcs else "entered manually."))
    derived = [l.label for l in inp.lines if l.source in ("derived", "approximation")]
    if derived or prof.derived:
        doc.add_paragraph("Derived or approximated lines: " + "; ".join(derived + prof.derived))
    d = doc.add_paragraph(DISCLAIMER)
    d.runs[0].font.size = Pt(9)
    d.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
